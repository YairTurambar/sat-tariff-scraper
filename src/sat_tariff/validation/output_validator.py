"""Workbook structure validation."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook
from ..exporters.layouts import FORBIDDEN_EXPORT_COLUMNS

EXPECTED_SHEETS = ["Derechos e impuestos", "Nomenclatura", "Restricciones", "Cuotas"]
POSITIONAL_HEADER_PATTERN = re.compile(r"^[A-Z0-9]+_TABLA(_\d+)?$", re.IGNORECASE)


class OutputValidationError(ValueError):
    pass



def validate_workbook_structure(path: str | Path) -> None:
    workbook = load_workbook(path)
    if workbook.sheetnames != EXPECTED_SHEETS:
        raise OutputValidationError(f"Unexpected sheet order: {workbook.sheetnames}")

    nomenclature = workbook["Nomenclatura"]
    merged_ranges = {str(cell_range) for cell_range in nomenclature.merged_cells.ranges}
    if not any(cell_range.startswith("I1:J1") for cell_range in merged_ranges):
        raise OutputValidationError("Nomenclatura sheet is missing the merged 'Unidades de medida' header.")

    rights_headers = [cell.value for cell in workbook["Derechos e impuestos"][1]]
    rights_row2 = [cell.value for cell in workbook["Derechos e impuestos"][2]]
    restrictions_headers = [cell.value for cell in workbook["Restricciones"][1]]
    quotas_headers = [cell.value for cell in workbook["Cuotas"][1]]
    nomenclature_headers = []
    for column_index in range(1, workbook["Nomenclatura"].max_column + 1):
        parent = workbook["Nomenclatura"].cell(row=1, column=column_index).value
        child = workbook["Nomenclatura"].cell(row=2, column=column_index).value
        nomenclature_headers.append(child or parent)
    if rights_headers[:4] != ["HS_Code", "Status", "Overall_Status", "Código"]:
        raise OutputValidationError("Unexpected Derechos e impuestos headers.")
    positional_headers = sorted(
        header
        for header in rights_headers
        if header and POSITIONAL_HEADER_PATTERN.match(str(header))
    )
    if positional_headers:
        raise OutputValidationError(
            "Derechos e impuestos leaks positional columns (table numbers never identify an "
            f"agreement): {positional_headers}"
        )
    if restrictions_headers[:4] != ["HS_Code", "Status", "Overall_Status", "Código"]:
        raise OutputValidationError("Unexpected Restricciones headers.")
    if quotas_headers[:4] != ["HS_Code", "Status", "Overall_Status", "TRATAMIENTO GENERAL"]:
        raise OutputValidationError("Unexpected Cuotas headers.")
    for sheet_name, headers in {
        "Derechos e impuestos": rights_headers + rights_row2,
        "Nomenclatura": nomenclature_headers,
        "Restricciones": restrictions_headers,
        "Cuotas": quotas_headers,
    }.items():
        leaked = FORBIDDEN_EXPORT_COLUMNS.intersection({header for header in headers if header})
        if leaked:
            raise OutputValidationError(f"{sheet_name} leaks forbidden headers: {sorted(leaked)}")
