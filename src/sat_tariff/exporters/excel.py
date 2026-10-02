"""Excel export implementation."""

from __future__ import annotations

from collections import OrderedDict
from datetime import datetime, timezone
import os
from pathlib import Path
import re
import shutil
from typing import Any
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from ..config import AppConfig
from ..validation.output_validator import validate_workbook_structure
from .layouts import DUTY_GROUP_ORDER_HINT, FORBIDDEN_EXPORT_COLUMNS, NOMENCLATURE_HEADERS, QUOTAS_HEADERS, RESTRICTIONS_HEADERS, RIGHTS_BASE_HEADERS, RIGHTS_TRAILING_HEADERS, SHEET_ORDER, TEXT_COLUMNS
from .styles import autosize_columns, style_data_cells, style_headers


def _ascii_slug(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value or "")
    ascii_value = normalized.encode("ascii", "ignore").decode("ascii").upper()
    slug = re.sub(r"[^A-Z0-9]+", "_", ascii_value).strip("_")
    return slug or "OTRO"



def _agreement_suffix(agreement_name: str) -> str:
    if agreement_name.strip().upper() == "TRATAMIENTO GENERAL":
        return "GENERAL"
    match = re.search(r"[-\u2013\u2014]\s*([A-Za-z]{2,8})\s*$", agreement_name)
    if match:
        return _ascii_slug(match.group(1))
    return _ascii_slug(agreement_name)



def _base_row(bundle: dict[str, Any], section_key: str) -> dict[str, Any]:
    return {
        "HS_Code": bundle["raw_code"],
        "Status": bundle.get(section_key, {}).get("status") or bundle["state"],
        "Overall_Status": bundle["state"],
    }



def _rights_column_metadata(bundles: list[dict[str, Any]]) -> "OrderedDict[str, tuple[str, str, str]]":
    """Map each dynamic rights column name to (code_label, suffix, group_label).

    ``group_label`` is the first full agreement name seen for that suffix
    (e.g. ``"TRATAMIENTO GENERAL"`` or ``"Tratado de Libre Comercio - MX"``),
    used as the merged group header above the per-code value columns.
    """
    metadata: "OrderedDict[str, tuple[str, str, str]]" = OrderedDict()
    for bundle in bundles:
        for row in bundle.get("rights", {}).get("rows", []):
            code_label = _ascii_slug(row.get("code", ""))
            agreement_name = row.get("agreement_name", "")
            suffix = _agreement_suffix(agreement_name)
            column_name = f"{code_label}_{suffix}"
            if column_name == "_":
                continue
            metadata.setdefault(column_name, (code_label, suffix, agreement_name.strip() or suffix))
    return metadata



def _rights_dynamic_columns(metadata: "OrderedDict[str, tuple[str, str, str]]") -> list[str]:
    ordered_names = list(metadata.keys())

    def sort_key(column_name: str) -> tuple[int, int, str]:
        _code, suffix, _label = metadata[column_name]
        if suffix == "GENERAL":
            return (0, ordered_names.index(column_name), column_name)
        hint_index = DUTY_GROUP_ORDER_HINT.index(suffix) if suffix in DUTY_GROUP_ORDER_HINT else len(DUTY_GROUP_ORDER_HINT)
        return (1 + hint_index, ordered_names.index(column_name), column_name)

    return sorted(ordered_names, key=sort_key)



def _build_rights_rows(bundles: list[dict[str, Any]], dynamic_columns: list[str]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for bundle in bundles:
        base = _base_row(bundle, "rights")
        section_rows = bundle.get("rights", {}).get("rows", [])
        if not section_rows:
            rows.append({**base, "Código": "", **{column: "" for column in dynamic_columns}, "Código adicional": "", "Código de cuota": ""})
            continue
        grouped: OrderedDict[tuple[str, str], dict[str, Any]] = OrderedDict()
        for source in section_rows:
            key = (source.get("additional_code", ""), source.get("quota_code", ""))
            target = grouped.setdefault(key, {**base, "Código": "", **{column: "" for column in dynamic_columns}, "Código adicional": key[0], "Código de cuota": key[1]})
            column_name = f"{_ascii_slug(source.get('code', ''))}_{_agreement_suffix(source.get('agreement_name', ''))}"
            existing = target.get(column_name, "")
            new_value = source.get("value", "")
            if new_value:
                target[column_name] = f"{existing}\n{new_value}".strip()
            source_code = source.get("code", "")
            if source_code:
                merged_codes = list(
                    OrderedDict.fromkeys(
                        filter(
                            None,
                            (target["Código"] + " | " + source_code).strip(" | ").split(" | "),
                        )
                    )
                )
                target["Código"] = " | ".join(merged_codes)
        rows.extend(grouped.values())
    return rows



def _build_nomenclature_rows(bundles: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for bundle in bundles:
        base = _base_row(bundle, "nomenclature")
        section_rows = bundle.get("nomenclature", {}).get("rows", [])
        scalars = {"Sección": "", "Capítulo:": "", "Fecha inicio de vigencia:": "", "Fecha fin de vigencia:": "", "Códigos adicionales": ""}
        units: list[tuple[str, str]] = []
        for record in section_rows:
            scalars["Sección"] = scalars["Sección"] or record.get("section", "")
            scalars["Capítulo:"] = scalars["Capítulo:"] or record.get("chapter", "")
            scalars["Fecha inicio de vigencia:"] = scalars["Fecha inicio de vigencia:"] or record.get("effective_from_raw", "")
            scalars["Fecha fin de vigencia:"] = scalars["Fecha fin de vigencia:"] or record.get("effective_to_raw", "")
            scalars["Códigos adicionales"] = scalars["Códigos adicionales"] or record.get("additional_codes_text", "")
            if record.get("record_type") == "unit":
                units.append((record.get("unit_code", ""), record.get("unit_description", "")))
        shared = {**base, **scalars}
        if not units:
            rows.append({**shared, "Código": "", "Descripción": ""})
        else:
            for code, description in units:
                rows.append({**shared, "Código": code, "Descripción": description})
    return rows


def _assert_no_forbidden_headers(headers: list[str]) -> None:
    leaked = FORBIDDEN_EXPORT_COLUMNS.intersection(headers)
    if leaked:
        raise ValueError(f"Forbidden export columns leaked into workbook: {sorted(leaked)}")



def _build_restrictions_rows(bundles: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for bundle in bundles:
        base = _base_row(bundle, "restrictions")
        section_rows = bundle.get("restrictions", {}).get("rows", [])
        if not section_rows:
            rows.append({**base, "Código": "", "Descripción": "", "Código adicional": "", "Valor": "", "Código de cuota": ""})
            continue
        for record in section_rows:
            rows.append({
                **base,
                "Código": record.get("code", ""),
                "Descripción": record.get("description", ""),
                "Código adicional": record.get("additional_code", ""),
                "Valor": record.get("value", ""),
                "Código de cuota": record.get("quota_code", ""),
            })
    return rows



def _build_quotas_rows(bundles: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for bundle in bundles:
        base = _base_row(bundle, "quotas")
        messages = [record.get("message", "") for record in bundle.get("quotas", {}).get("rows", []) if record.get("message")]
        rows.append({**base, "TRATAMIENTO GENERAL": "\n".join(messages)})
    return rows



def _write_headers(ws, headers: list[str]) -> None:
    ws.append(headers)



def _write_rights_sheet(
    ws,
    rows: list[dict[str, Any]],
    dynamic_columns: list[str],
    column_metadata: "OrderedDict[str, tuple[str, str, str]]",
) -> list[str]:
    """Write the two-row grouped header (agreement group over duty-code columns).

    Row 1 shows the agreement name (e.g. ``TRATAMIENTO GENERAL`` or a treaty
    name) merged across every value column that belongs to it. Row 2 shows the
    specific duty/tax code (e.g. ``DAI``, ``IVA``) for each value column.
    Fixed columns (``HS_Code``, ``Código adicional``, ...) are merged
    vertically across both header rows, matching the ``Nomenclatura`` sheet's
    layout convention for non-grouped columns.
    """
    all_headers = RIGHTS_BASE_HEADERS + dynamic_columns + RIGHTS_TRAILING_HEADERS
    ws.append([None] * len(all_headers))
    row2_values = (
        list(RIGHTS_BASE_HEADERS)
        + [column_metadata[column_name][0] for column_name in dynamic_columns]
        + list(RIGHTS_TRAILING_HEADERS)
    )
    ws.append(row2_values)

    column_index = 1
    for header in RIGHTS_BASE_HEADERS:
        ws.cell(row=1, column=column_index, value=header)
        ws.merge_cells(start_row=1, start_column=column_index, end_row=2, end_column=column_index)
        column_index += 1

    group_spans: "OrderedDict[str, list[int]]" = OrderedDict()
    group_labels: dict[str, str] = {}
    for column_name in dynamic_columns:
        _code_label, suffix, label = column_metadata[column_name]
        group_spans.setdefault(suffix, []).append(column_index)
        group_labels.setdefault(suffix, label)
        column_index += 1
    for suffix, indices in group_spans.items():
        start_column, end_column = min(indices), max(indices)
        ws.cell(row=1, column=start_column, value=group_labels[suffix])
        if start_column != end_column:
            ws.merge_cells(start_row=1, start_column=start_column, end_row=1, end_column=end_column)

    for header in RIGHTS_TRAILING_HEADERS:
        ws.cell(row=1, column=column_index, value=header)
        ws.merge_cells(start_row=1, start_column=column_index, end_row=2, end_column=column_index)
        column_index += 1

    for row in rows:
        ws.append([row.get(header, "") for header in all_headers])

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(max(1, len(all_headers)))}{max(ws.max_row, 2)}"
    return all_headers



def _write_nomenclature_sheet(ws, rows: list[dict[str, Any]]) -> None:
    ws.append(NOMENCLATURE_HEADERS[:-2] + ["Unidades de medida", None])
    ws.append(NOMENCLATURE_HEADERS[:-2] + ["Código", "Descripción"])
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)
    ws.merge_cells(start_row=1, start_column=3, end_row=2, end_column=3)
    ws.merge_cells(start_row=1, start_column=4, end_row=2, end_column=4)
    ws.merge_cells(start_row=1, start_column=5, end_row=2, end_column=5)
    ws.merge_cells(start_row=1, start_column=6, end_row=2, end_column=6)
    ws.merge_cells(start_row=1, start_column=7, end_row=2, end_column=7)
    ws.merge_cells(start_row=1, start_column=8, end_row=2, end_column=8)
    ws.merge_cells(start_row=1, start_column=9, end_row=1, end_column=10)
    for row in rows:
        ws.append([row.get(header, "") for header in NOMENCLATURE_HEADERS])
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:J{max(ws.max_row, 2)}"
    style_headers(ws, 2)
    style_data_cells(ws, 3)



def _apply_text_formats(ws, header_rows: int) -> None:
    header_values = {}
    for row_index in range(1, min(ws.max_row, header_rows) + 1):
        for col_index in range(1, ws.max_column + 1):
            value = ws.cell(row=row_index, column=col_index).value
            if value:
                header_values[col_index] = value
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            header = header_values.get(cell.column)
            if header in TEXT_COLUMNS:
                cell.number_format = "@"



def _backup_existing(destination: Path) -> None:
    backup_dir = destination.parent / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")
    shutil.copy2(destination, backup_dir / f"{destination.stem}-{timestamp}{destination.suffix}")



def export_workbook(bundles: list[dict[str, Any]], config: AppConfig, output_path: str | Path | None = None) -> Path:
    destination = Path(output_path or config.output_xlsx)
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists() and config.backup_output:
        _backup_existing(destination)

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    dynamic_rights_metadata = _rights_column_metadata(bundles)
    dynamic_rights = _rights_dynamic_columns(dynamic_rights_metadata)
    rights_headers = RIGHTS_BASE_HEADERS + dynamic_rights + RIGHTS_TRAILING_HEADERS
    sections = {
        "Derechos e impuestos": _build_rights_rows(bundles, dynamic_rights),
        "Nomenclatura": _build_nomenclature_rows(bundles),
        "Restricciones": _build_restrictions_rows(bundles),
        "Cuotas": _build_quotas_rows(bundles),
    }

    for sheet_name in SHEET_ORDER:
        ws = workbook.create_sheet(sheet_name)
        if sheet_name == "Derechos e impuestos":
            _assert_no_forbidden_headers(rights_headers)
            _write_rights_sheet(ws, sections[sheet_name], dynamic_rights, dynamic_rights_metadata)
            style_headers(ws, 2)
            style_data_cells(ws, 3)
            header_rows = 2
        elif sheet_name == "Nomenclatura":
            _assert_no_forbidden_headers(NOMENCLATURE_HEADERS)
            _write_nomenclature_sheet(ws, sections[sheet_name])
            header_rows = 2
        elif sheet_name == "Restricciones":
            _assert_no_forbidden_headers(RESTRICTIONS_HEADERS)
            _write_headers(ws, RESTRICTIONS_HEADERS)
            for row in sections[sheet_name]:
                ws.append([row.get(header, "") for header in RESTRICTIONS_HEADERS])
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:H{max(ws.max_row, 1)}"
            style_headers(ws, 1)
            style_data_cells(ws, 2)
            header_rows = 1
        else:
            _assert_no_forbidden_headers(QUOTAS_HEADERS)
            _write_headers(ws, QUOTAS_HEADERS)
            for row in sections[sheet_name]:
                ws.append([row.get(header, "") for header in QUOTAS_HEADERS])
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:D{max(ws.max_row, 1)}"
            style_headers(ws, 1)
            style_data_cells(ws, 2)
            header_rows = 1
        autosize_columns(ws)
        _apply_text_formats(ws, header_rows)

    temp_path = destination.with_name(f"{destination.stem}.tmp{destination.suffix}")
    workbook.save(temp_path)
    load_workbook(temp_path)
    validate_workbook_structure(temp_path)
    os.replace(temp_path, destination)
    return destination
