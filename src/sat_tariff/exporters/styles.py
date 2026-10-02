"""Workbook style helpers."""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(color="FFFFFF", bold=True)
HEADER_ROW_HEIGHT = 22
#: Explicit height (in points) reserved for one line of text in a data row.
#: Data rows must not rely on the application's automatic row height: Excel and
#: LibreOffice use different defaults (15 pt vs ~12.8 pt) and different
#: wrapped-text metrics, so an implicit height makes the sheet render
#: differently in each engine. Pinning it keeps the output reproducible and
#: matches the row proportions of the reference screenshots.
DATA_LINE_HEIGHT = 25
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)



def style_headers(worksheet, header_rows: int) -> None:
    for row_index in range(1, header_rows + 1):
        worksheet.row_dimensions[row_index].height = HEADER_ROW_HEIGHT
        for cell in worksheet[row_index]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)



def _cell_line_count(cell, column_width: float) -> int:
    """Estimate how many rendered lines *cell* needs at *column_width* chars."""
    text = "" if cell.value is None else str(cell.value)
    if not text:
        return 1
    usable_width = max(1.0, column_width - 1)
    lines = 0
    for segment in text.split("\n"):
        lines += max(1, -(-len(segment) // int(usable_width)))
    return max(1, lines)



def style_data_cells(worksheet, start_row: int) -> None:
    for row in worksheet.iter_rows(min_row=start_row):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)



def set_data_row_heights(worksheet, start_row: int) -> None:
    """Pin an explicit height on every data row.

    Must run *after* :func:`autosize_columns`, because the estimated number of
    wrapped lines depends on the final column widths.
    """
    for row in worksheet.iter_rows(min_row=start_row):
        if not row:
            continue
        lines = 1
        for cell in row:
            width = worksheet.column_dimensions[get_column_letter(cell.column)].width or 8.43
            lines = max(lines, _cell_line_count(cell, width))
        worksheet.row_dimensions[row[0].row].height = DATA_LINE_HEIGHT * lines



def autosize_columns(worksheet) -> None:
    for index in range(1, worksheet.max_column + 1):
        width = max(len(str(worksheet.cell(row=row, column=index).value or "")) for row in range(1, worksheet.max_row + 1))
        worksheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 12), 60)
