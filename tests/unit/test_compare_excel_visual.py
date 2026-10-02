"""Unit tests for the pure, tool-independent helpers in
``scripts/compare_excel_visual.py``.

These exercise logic that does not require LibreOffice/poppler or the
optional ``visual`` extra to be installed, so they always run as part of the
default (structural) test suite -- unlike ``tests/visual/test_visual_regression.py``,
which is gated on those external tools/dependencies/reference files.
"""

from __future__ import annotations

from pathlib import Path
import sys

REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT / "scripts") not in sys.path:
    sys.path.insert(0, str(REPO_ROOT / "scripts"))

from compare_excel_visual import (  # noqa: E402
    REAL_REFERENCE_FILE_ALIASES,
    SHEET_REFERENCE_FILES,
    _prepare_print_ready_copy,
    missing_reference_files,
    missing_render_tools,
    missing_python_deps,
    resolve_reference_path,
)
from sample_workbook import export_sample_workbook  # noqa: E402


def test_missing_render_tools_returns_list_of_strings():
    missing = missing_render_tools()
    assert isinstance(missing, list)
    assert all(isinstance(tool, str) for tool in missing)
    assert set(missing) <= {"soffice", "pdftoppm"}


def test_missing_python_deps_returns_list_of_strings():
    missing = missing_python_deps()
    assert isinstance(missing, list)
    assert all(isinstance(dep, str) for dep in missing)
    assert set(missing) <= {"PIL", "numpy", "skimage"}


def test_missing_reference_files_reports_all_four_when_dir_is_empty(tmp_path):
    missing = missing_reference_files(tmp_path)
    assert sorted(missing) == sorted(SHEET_REFERENCE_FILES.values())


def test_missing_reference_files_reports_only_absent_files(tmp_path):
    present = next(iter(SHEET_REFERENCE_FILES.values()))
    (tmp_path / present).write_bytes(b"not-a-real-png")

    missing = missing_reference_files(tmp_path)

    assert present not in missing
    assert set(missing) == set(SHEET_REFERENCE_FILES.values()) - {present}


def test_missing_reference_files_empty_when_all_present(tmp_path):
    for filename in SHEET_REFERENCE_FILES.values():
        (tmp_path / filename).write_bytes(b"not-a-real-png")

    assert missing_reference_files(tmp_path) == []


def test_real_screenshot_names_resolve_to_canonical_references(tmp_path):
    for filename in REAL_REFERENCE_FILE_ALIASES.values():
        (tmp_path / filename).write_bytes(b"real-reference-placeholder")

    assert missing_reference_files(tmp_path) == []
    for canonical, real_name in REAL_REFERENCE_FILE_ALIASES.items():
        assert resolve_reference_path(tmp_path, canonical) == tmp_path / real_name


def test_prepare_print_ready_copy_sets_fit_to_page_for_every_sheet(tmp_path):
    import zipfile

    source = export_sample_workbook(tmp_path / "sample.xlsx")
    destination = tmp_path / "sample.print-ready.xlsx"

    _prepare_print_ready_copy(source, destination)

    # Inspect the written OOXML directly: reloading page_setup with openpyxl
    # after a save/reload round trip is unreliable for this attribute, so we
    # assert on the persisted XML instead of re-parsing with openpyxl.
    with zipfile.ZipFile(destination) as archive:
        sheet_xml_names = [
            name for name in archive.namelist() if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
        ]
        assert sheet_xml_names, "expected at least one worksheet in the print-ready copy"
        for name in sheet_xml_names:
            xml = archive.read(name).decode("utf-8")
            assert 'fitToPage="1"' in xml
            assert 'orientation="landscape"' in xml


def _write_manifest(tmp_path, overrides=None, image_bytes=b"baseline-png"):
    import json

    from compare_excel_visual import RENDER_DPI, SHEET_ORDER, sha256_of

    sheets = {}
    for sheet_name in SHEET_ORDER:
        filename = SHEET_REFERENCE_FILES[sheet_name]
        (tmp_path / filename).write_bytes(image_bytes)
        sheets[sheet_name] = {
            "file": filename,
            "dimensions": [10, 10],
            "sha256": sha256_of(tmp_path / filename),
        }
    manifest = {"dpi": RENDER_DPI, "sheets": sheets}
    manifest.update(overrides or {})
    (tmp_path / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    return manifest


def test_load_baseline_manifest_accepts_a_complete_baseline(tmp_path):
    from compare_excel_visual import load_baseline_manifest

    _write_manifest(tmp_path)

    manifest = load_baseline_manifest(tmp_path)

    assert set(manifest["sheets"]) == set(SHEET_REFERENCE_FILES)


def test_load_baseline_manifest_reports_a_missing_baseline(tmp_path):
    import pytest

    from compare_excel_visual import BaselineUnavailable, load_baseline_manifest

    with pytest.raises(BaselineUnavailable) as error:
        load_baseline_manifest(tmp_path)

    assert "update_visual_baseline.py --confirm" in str(error.value)


def test_load_baseline_manifest_rejects_a_different_render_dpi(tmp_path):
    import pytest

    from compare_excel_visual import RENDER_DPI, BaselineUnavailable, load_baseline_manifest

    _write_manifest(tmp_path, overrides={"dpi": RENDER_DPI + 50})

    with pytest.raises(BaselineUnavailable) as error:
        load_baseline_manifest(tmp_path)

    assert "dpi" in str(error.value)


def test_load_baseline_manifest_rejects_a_tampered_image(tmp_path):
    import pytest

    from compare_excel_visual import BaselineUnavailable, load_baseline_manifest

    _write_manifest(tmp_path)
    (tmp_path / SHEET_REFERENCE_FILES["Cuotas"]).write_bytes(b"tampered")

    with pytest.raises(BaselineUnavailable) as error:
        load_baseline_manifest(tmp_path)

    assert "does not match its manifest hash" in str(error.value)


def test_update_visual_baseline_requires_explicit_confirmation(tmp_path):
    from update_visual_baseline import main

    exit_code = main(["--baseline-dir", str(tmp_path)])

    assert exit_code == 2
    assert list(tmp_path.iterdir()) == []


def test_committed_baseline_manifest_is_consistent():
    from compare_excel_visual import BASELINE_DIR, load_baseline_manifest

    if not (BASELINE_DIR / "manifest.json").is_file():
        import pytest

        pytest.skip("No renderer baseline committed in this checkout.")

    manifest = load_baseline_manifest(BASELINE_DIR)

    assert manifest["dpi"] == 200
    assert manifest["libreoffice_version"]
    assert manifest["generated_at"]
    assert manifest["command"].endswith("--confirm")


def test_design_mode_validates_workbook_structure(tmp_path):
    from compare_excel_visual import validate_workbook_against_spec

    workbook_path = export_sample_workbook(tmp_path / "sample.xlsx")

    validate_workbook_against_spec(workbook_path)


def test_structure_validation_detects_a_header_color_regression(tmp_path):
    import pytest
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

    from compare_excel_visual import validate_workbook_against_spec

    workbook_path = export_sample_workbook(tmp_path / "sample.xlsx")
    workbook = load_workbook(workbook_path)
    workbook["Derechos e impuestos"]["A1"].fill = PatternFill("solid", fgColor="FF0000")
    workbook.save(workbook_path)

    with pytest.raises(ValueError, match="Header fill mismatch"):
        validate_workbook_against_spec(workbook_path)
