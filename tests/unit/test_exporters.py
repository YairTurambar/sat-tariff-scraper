from pathlib import Path
import os
import subprocess
import sys
import uuid

from openpyxl import load_workbook
import pytest

from sat_tariff.config import AppConfig
from sat_tariff.exporters.layouts import FORBIDDEN_EXPORT_COLUMNS
from sat_tariff.exporters.excel import export_workbook
from sat_tariff.exporters.styles import DATA_LINE_HEIGHT
from sat_tariff.storage import Storage


def test_corrected_sections_round_trip_sqlite_to_excel(tmp_path):
    from sat_tariff.extractors.nomenclature import parse_nomenclature
    from sat_tariff.extractors.quotas import parse_quotas
    from sat_tariff.extractors.restrictions import NO_RESTRICTIONS_MESSAGE, parse_restrictions
    from sat_tariff.models import ProcessingState
    from sat_tariff.validation.output_validator import OutputValidationError, validate_workbook_structure

    fixtures = Path(__file__).resolve().parents[1] / "fixtures"
    storage = Storage(tmp_path / "source.sqlite3")
    try:
        storage.upsert_code("0210110000", "0210110000", ProcessingState.completed)
        for section, parser, fixture in (
            ("nomenclature", parse_nomenclature, "nomenclature_caption.html"),
            ("restrictions", parse_restrictions, "restrictions_no_permits.html"),
            ("quotas", parse_quotas, "quotas_structured.html"),
        ):
            result = parser((fixtures / fixture).read_text(encoding="utf-8"))
            storage.save_section_rows("0210110000", section, result.rows, section_status=result.status)
        # A historical JSON record must not require a schema migration.
        storage.upsert_code("0101210000", "0101210000", ProcessingState.completed)
        storage.save_section_rows("0101210000", "quotas", [{
            "treatment_name": "TRATAMIENTO GENERAL", "message": "Cupo anual | 100 toneladas", "source_status": "ok",
        }], section_status="ok")
        bundles = storage.load_all_for_export()
    finally:
        storage.close()
    path = export_workbook(bundles, AppConfig(output_xlsx=tmp_path / "export.xlsx", backup_output=False))
    workbook = load_workbook(path)
    quotas = workbook["Cuotas"]
    assert [cell.value for cell in quotas[1]] == [
        "HS_Code", "Status", "Overall_Status", "TRATAMIENTO GENERAL", "Cuota", "País",
        "Inicio vigencia", "Fin vigencia", "Cantidad Inicial Permitida", "Fecha Agotamiento", "% Variación",
    ]
    assert quotas.max_row == 4
    assert quotas["D2"].value == "Cupo anual | 100 toneladas"
    assert all(quotas.cell(2, column).value is None for column in range(5, 12))
    assert [cell.value for cell in quotas[3]][4:] == [
        "Cuota 1", "Guatemala", "01/01/2024", "31/12/2024", "100", "30/11/2024", "5%",
    ]
    assert [cell.value for cell in quotas[4]][4:] == ["Cuota 2", None, "01/01/2025", None, "200", None, "0%"]
    assert quotas["D3"].value == quotas["D4"].value
    assert quotas.auto_filter.ref == "A1:K4"
    restrictions = workbook["Restricciones"]
    assert [cell.value for cell in restrictions[1]] == [
        "HS_Code", "Status", "Overall_Status", "Código", "Descripción", "Código adicional", "Valor", "Código de cuota",
    ]
    assert [cell.value for cell in restrictions[3]] == [
        "0210110000", "ok", "completed", None, NO_RESTRICTIONS_MESSAGE, None, None, None,
    ]
    nomenclature = workbook["Nomenclatura"]
    assert [nomenclature.cell(2, i).value or nomenclature.cell(1, i).value for i in range(1, 11)] == [
        "HS_Code", "Status", "Overall_Status", "Sección", "Capítulo:", "Fecha inicio de vigencia:",
        "Fecha fin de vigencia:", "Códigos adicionales", "Código", "Descripción",
    ]
    assert nomenclature["I4"].value == "0210.11.0000"
    assert nomenclature["J4"].value == "- - Jamones, paletas y sus trozos, sin deshuesar"
    assert nomenclature["I5"].value == "KGM"
    assert nomenclature["J5"].value == "Kilogramo"
    assert nomenclature["D4"].value == nomenclature["D5"].value == "Sección I"
    for name in ("Derechos e impuestos", "Restricciones", "Nomenclatura"):
        sheet = workbook[name]
        assert sheet["A1"].fill.fgColor.rgb[-6:] == "4472C4"
        assert sheet["A1"].font.bold
        assert sheet["A1"].font.color.rgb[-6:] == "FFFFFF"
        assert sheet.freeze_panes == ("A3" if name == "Nomenclatura" else "A2")
        assert sheet.cell(3 if name == "Nomenclatura" else 2, 1).number_format == "@"
    assert restrictions.auto_filter.ref == "A1:H3"
    assert nomenclature.auto_filter.ref == "A2:J5"
    validate_workbook_structure(path)
    quotas.delete_cols(11)
    workbook.save(tmp_path / "invalid.xlsx")
    workbook.close()
    with pytest.raises(OutputValidationError, match="Cuotas"):
        validate_workbook_structure(tmp_path / "invalid.xlsx")


def test_nomenclature_historical_summary_and_other_details_are_not_lost():
    from sat_tariff.exporters.excel import _build_nomenclature_rows

    rows = _build_nomenclature_rows([{
        "raw_code": "0210110000", "state": "completed",
        "nomenclature": {"status": "empty", "rows": [
            {"record_type": "summary", "section": "Sección I", "chapter": "Capítulo 02"},
            {"record_type": "goods", "code": "0210.11.0000", "description": "Jamones"},
            {"record_type": "unit", "unit_code": "KGM", "unit_description": "Kilogramo"},
            {"record_type": "unit", "unit_code": "KGM", "unit_description": "Kilogramo"},
            {"record_type": "Clasificadores estadísticos", "content": "No aplica"},
        ]},
    }])
    assert [(row["Código"], row["Descripción"]) for row in rows] == [
        ("0210.11.0000", "Jamones"), ("KGM", "Kilogramo"), ("", "No aplica"),
    ]
    assert all(row["Status"] == "ok" and row["Sección"] == "Sección I" for row in rows)
    partial = _build_nomenclature_rows([{
        "raw_code": "0210110000", "state": "completed",
        "nomenclature": {"status": "empty", "rows": [{"record_type": "summary", "section": "Sección I"}]},
    }])
    assert len(partial) == 1
    assert partial[0]["Status"] == "ok"
    assert partial[0]["Sección"] == "Sección I"


def test_quotas_do_not_overwrite_shared_treatment_with_null():
    from sat_tariff.exporters.excel import _build_quotas_rows

    rows = _build_quotas_rows([{
        "raw_code": "0210110000", "state": "completed", "quotas": {"status": "ok", "rows": [
            {"general_treatment": "0210.11.0000 | Jamones", "quota": "1"},
            {"general_treatment": None, "quota": "2"},
            {"general_treatment": "", "quota": "3"},
        ]},
    }])
    assert len(rows) == 3
    assert all(row["TRATAMIENTO GENERAL"] == "0210.11.0000 | Jamones" for row in rows)


@pytest.mark.parametrize("first_import", [
    "import sat_tariff.validation.output_validator",
    "import sat_tariff.extractors.quotas",
])
def test_direct_validator_and_package_export_imports_in_clean_process(first_import):
    root = Path(__file__).resolve().parents[2]
    env = {**os.environ, "PYTHONPATH": str(root / "src")}
    subprocess.run([
        sys.executable, "-c",
        first_import + "; from sat_tariff.exporters import export_workbook; "
        "import sys; assert callable(export_workbook); assert 'playwright.sync_api' not in sys.modules",
    ], check=True, env=env, cwd=root)


def make_paths():
    suffix = uuid.uuid4().hex
    db_path = Path("data") / f"export-test-{suffix}.sqlite3"
    xlsx_path = Path("data") / f"export-test-{suffix}.xlsx"
    for path in (db_path, xlsx_path):
        path.unlink(missing_ok=True)
    return db_path, xlsx_path



def test_exporter_builds_expected_workbook_structure():
    db_path, xlsx_path = make_paths()
    storage = Storage(db_path)
    try:
        storage.upsert_code("9999000001", "9999000001", last_error=None)
        storage.save_section_rows(
            "9999000001",
            "rights",
            [
                {
                    "agreement_name": "TRATAMIENTO GENERAL",
                    "code": "DAI",
                    "description": "General",
                    "additional_code": "AD1",
                    "value": "5%",
                    "quota_code": "CQ1",
                },
                {
                    "agreement_name": "Tratado de Libre Comercio - MX",
                    "code": "DAI",
                    "description": "Preferencial",
                    "additional_code": "AD1",
                    "value": "0%",
                    "quota_code": "CQ1",
                },
            ],
            section_status="ok",
        )
        storage.save_section_rows(
            "9999000001",
            "nomenclature",
            [
                {
                    "record_type": "unit",
                    "section": "Sección I",
                    "chapter": "Capítulo 01",
                    "effective_from_raw": "01/01/2024",
                    "effective_to_raw": "31/12/2024",
                    "additional_codes_text": "No se han encontrado códigos adicionales asociados al inciso consultado",
                    "unit_code": "KGM",
                    "unit_description": "Kilogramo",
                }
            ],
            section_status="ok",
        )
        storage.save_section_rows(
            "9999000001",
            "restrictions",
            [{"code": "R1", "description": "Licencia", "additional_code": "AD2", "value": "Sí", "quota_code": "CQ2"}],
            section_status="ok",
        )
        storage.save_section_rows(
            "9999000001",
            "quotas",
            [{"message": "No se han encontrado cuotas/contingentes para el inciso consultado"}],
            section_status="no_quotas",
        )
        bundles = storage.load_all_for_export()
    finally:
        storage.close()

    config = AppConfig(sqlite_db=db_path, output_xlsx=xlsx_path, backup_output=False)
    path = export_workbook(bundles, config)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Derechos e impuestos", "Nomenclatura", "Restricciones", "Cuotas"]
    assert "Sheet" not in workbook.sheetnames

    rights = workbook["Derechos e impuestos"]
    nomenclature = workbook["Nomenclatura"]
    restrictions = workbook["Restricciones"]
    quotas = workbook["Cuotas"]

    assert rights["A1"].value == "HS_Code"
    assert rights["A1"].fill.fgColor.rgb[-6:] == "4472C4"
    assert rights["A1"].font.bold is True
    assert rights["A1"].font.color.rgb[-6:] == "FFFFFF"
    assert rights.row_dimensions[1].height == 22
    # Data rows carry an explicit height so Excel and LibreOffice (whose
    # automatic row heights differ) render the same geometry.
    assert rights.row_dimensions[2].height == DATA_LINE_HEIGHT
    assert rights.freeze_panes == "A2"
    assert rights["A2"].number_format == "@"
    rights_row1 = [cell.value for cell in rights[1]]
    assert rights_row1[4:6] == ["DAI_GENERAL", "DAI_MX"]

    assert any(str(cell_range) == "I1:J1" for cell_range in nomenclature.merged_cells.ranges)
    assert nomenclature.freeze_panes == "A3"
    assert nomenclature["A3"].number_format == "@"

    assert restrictions["A1"].value == "HS_Code"
    assert quotas["D1"].value == "TRATAMIENTO GENERAL"
    for sheet in (restrictions, quotas):
        headers = [cell.value for cell in sheet[1] if cell.value]
        assert FORBIDDEN_EXPORT_COLUMNS.isdisjoint(headers)
    rights_headers = [value for value in rights_row1 if value]
    assert FORBIDDEN_EXPORT_COLUMNS.isdisjoint(rights_headers)
    nomenclature_headers = [
        nomenclature.cell(row=2, column=index).value or nomenclature.cell(row=1, column=index).value
        for index in range(1, nomenclature.max_column + 1)
    ]
    assert FORBIDDEN_EXPORT_COLUMNS.isdisjoint({header for header in nomenclature_headers if header})

    workbook.close()
    path.unlink(missing_ok=True)
    db_path.unlink(missing_ok=True)


def test_rights_sheet_uses_flat_dynamic_headers(tmp_path):
    import sys

    scripts_dir = Path(__file__).resolve().parents[2] / "scripts"
    if str(scripts_dir) not in sys.path:
        sys.path.insert(0, str(scripts_dir))
    from sample_workbook import export_sample_workbook

    path = export_sample_workbook(tmp_path / "sample.xlsx")
    workbook = load_workbook(path)
    rights = workbook["Derechos e impuestos"]

    row1 = [cell.value for cell in rights[1]]
    assert row1[4:6] == ["DAI_GENERAL", "IVA_GENERAL"]
    assert row1[6] == "DAI_CL"

    workbook.close()


def test_rights_sheet_keeps_values_with_their_agreements_across_codes(tmp_path):
    mexico = (
        "Tratado de Libre Comercio Entre Los Estados Unidos Mexicanos y las "
        "Repúblicas de Costa Rica, El Salvador, Guatemala, Honduras y Nicaragua - MX"
    )
    mexico_alias = (
        "TRATADO DE LIBRE COMERCIO ENTRE LOS ESTADOS UNIDOS MEXICANOS Y LAS "
        "REPUBLICAS DE COSTA RICA, EL SALVADOR, GUATEMALA, HONDURAS Y NICARAGUA — MX"
    )
    belize = (
        "Acuerdo de Alcance Parcial entre el Gobierno de la República de Guatemala "
        "y el Gobierno de Belice - BZ"
    )
    bundles = [
        {
            "raw_code": "1111111111",
            "state": "completed",
            "rights": {
                "status": "ok",
                "rows": [
                    {"agreement_name": belize, "code": "DAI", "value": "2%"},
                    {"agreement_name": mexico, "code": "DAI", "value": "7%"},
                ],
            },
            "nomenclature": {"rows": []},
            "restrictions": {"rows": []},
            "quotas": {"rows": []},
        },
        {
            "raw_code": "2222222222",
            "state": "completed",
            "rights": {
                "status": "ok",
                "rows": [
                    {"agreement_name": "Nuevo Acuerdo Comercial - XY", "code": "DAI", "value": "1%"},
                    {"agreement_name": mexico_alias, "code": "DAI", "value": "3%"},
                ],
            },
            "nomenclature": {"rows": []},
            "restrictions": {"rows": []},
            "quotas": {"rows": []},
        },
    ]
    path = tmp_path / "agreements.xlsx"
    config = AppConfig(output_xlsx=path, backup_output=False)

    export_workbook(bundles, config)

    workbook = load_workbook(path)
    rights = workbook["Derechos e impuestos"]
    headers = [cell.value for cell in rights[1]]
    assert headers.count("DAI_MX") == 1
    rows_by_code = {
        rights.cell(row=row_index, column=1).value: {
            header: rights.cell(row=row_index, column=column_index).value
            for column_index, header in enumerate(headers, start=1)
        }
        for row_index in range(2, rights.max_row + 1)
    }
    assert rows_by_code["1111111111"]["DAI_MX"] == "7%"
    assert rows_by_code["1111111111"]["DAI_BZ"] == "2%"
    assert rows_by_code["1111111111"]["DAI_XY"] is None
    assert rows_by_code["2222222222"]["DAI_MX"] == "3%"
    assert rows_by_code["2222222222"]["DAI_BZ"] is None
    assert rows_by_code["2222222222"]["DAI_XY"] == "1%"
    workbook.close()


def test_rights_sheet_warns_and_separates_unknown_agreements(tmp_path, caplog):
    bundles = [
        {
            "raw_code": "3333333333",
            "state": "completed",
            "rights": {
                "status": "ok",
                "rows": [
                    {"agreement_name": "Acuerdo Uno - XY", "code": "DAI", "value": "1%"},
                    {"agreement_name": "Acuerdo Dos - XY", "code": "DAI", "value": "2%"},
                    {"agreement_name": "Acuerdo sin identificador", "code": "DAI", "value": "3%"},
                    {"agreement_name": "Acuerdo Uno - XY", "code": "DAI", "value": "1%"},
                ],
            },
            "nomenclature": {"rows": []},
            "restrictions": {"rows": []},
            "quotas": {"rows": []},
        }
    ]
    path = tmp_path / "unknown-agreements.xlsx"
    export_workbook(bundles, AppConfig(output_xlsx=path, backup_output=False))

    workbook = load_workbook(path)
    rights = workbook["Derechos e impuestos"]
    headers = [cell.value for cell in rights[1]]
    assert headers[4:6] == ["DAI_XY", "DAI_XY_2"]
    values = {
        headers[column - 1]: rights.cell(row=2, column=column).value
        for column in range(1, rights.max_column + 1)
    }
    assert {values["DAI_XY"], values["DAI_XY_2"]} == {"1%", "2%"}
    assert "Acuerdo sin identificador" in caplog.text
    assert "DAI_ACUERDO" not in headers
    workbook.close()


def test_rights_sheet_from_clean_extraction_uses_semantic_headers(tmp_path):
    from sat_tariff.extractors.rights_taxes import SAT_AGREEMENT_COLUMN_MAP, parse_rights_taxes

    html = Path("tests/fixtures/rights_taxes_jsf.html").read_text(encoding="utf-8")
    result = parse_rights_taxes(html)
    bundles = [
        {
            "raw_code": "4444444444",
            "state": "completed",
            "rights": {"status": result.status, "rows": [row.to_dict() for row in result.rows]},
            "nomenclature": {"rows": []},
            "restrictions": {"rows": []},
            "quotas": {"rows": []},
        }
    ]
    path = tmp_path / "clean.xlsx"
    export_workbook(bundles, AppConfig(output_xlsx=path, backup_output=False))

    workbook = load_workbook(path)
    headers = [cell.value for cell in workbook["Derechos e impuestos"][1]]
    assert headers[4:6] == ["DAI_GENERAL", "IVA_GENERAL"]
    assert set(SAT_AGREEMENT_COLUMN_MAP.values()).issubset(set(headers))
    assert not any(str(header).startswith(("DAI_TABLA", "IVA_TABLA")) for header in headers if header)
    workbook.close()


def test_rights_sheet_omits_positional_identities_and_reports_them(tmp_path, caplog):
    from sat_tariff.exporters.excel import collect_positional_rights_issues

    bundles = [
        {
            "raw_code": "5555555555",
            "state": "completed",
            "rights": {
                "status": "ok",
                "rows": [
                    {"agreement_name": "Tabla 14", "code": "DAI", "value": "10%"},
                    {"agreement_name": "Tabla 14", "code": "IVA", "value": "12%"},
                    {"agreement_name": "Tabla 17", "code": "DAI", "value": "0%"},
                    {"agreement_name": "TRATAMIENTO GENERAL", "code": "DAI", "value": "5%"},
                ],
            },
            "nomenclature": {"rows": []},
            "restrictions": {"rows": []},
            "quotas": {"rows": []},
        }
    ]
    assert collect_positional_rights_issues(bundles) == {"5555555555": ["Tabla 14", "Tabla 17"]}

    path = tmp_path / "positional.xlsx"
    export_workbook(bundles, AppConfig(output_xlsx=path, backup_output=False))

    workbook = load_workbook(path)
    headers = [cell.value for cell in workbook["Derechos e impuestos"][1] if cell.value]
    assert "DAI_GENERAL" in headers
    assert not any(str(header).startswith(("DAI_TABLA", "IVA_TABLA")) for header in headers)
    assert "repair-rights" in caplog.text
    workbook.close()
