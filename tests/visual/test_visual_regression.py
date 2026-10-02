"""Visual-regression tests for the exported SAT workbook (dual reference).

Two independent levels are exercised here:

1. **Design reference** (``references/*.png``): the original Excel screenshots.
   LibreOffice is not Excel, so this comparison is tolerant and is paired with
   exact structural assertions made through ``openpyxl``
   (``validate_workbook_against_spec`` plus ``tests/unit/test_exporters.py``).
2. **Renderer baseline** (``tests/visual/baselines/libreoffice/``): PNGs
   produced by this very pipeline from the approved workbook. Both sides come
   from the same engine, so the thresholds are strict and the test also proves
   that a material change to the workbook is detected.

Each test self-skips, with an explicit human-readable reason, whenever a
precondition is missing (LibreOffice/poppler, the optional ``visual`` extra,
the reference screenshots or an approved baseline). Tests never regenerate the
baseline: use ``python scripts/update_visual_baseline.py --confirm``.
"""

from __future__ import annotations

from pathlib import Path
import sys

import pytest

REPO_ROOT = Path(__file__).resolve().parents[2]
SCRIPTS_DIR = REPO_ROOT / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

from compare_excel_visual import (  # noqa: E402
    BASELINE_DIR,
    BASELINE_MANIFEST_NAME,
    BASELINE_MEAN_DIFF_THRESHOLD,
    BASELINE_SSIM_THRESHOLD,
    DESIGN_MEAN_DIFF_THRESHOLD,
    DESIGN_SSIM_THRESHOLD,
    SHEET_REFERENCE_FILES,
    BaselineUnavailable,
    compare_workbook_to_baseline,
    compare_workbook_to_references,
    missing_python_deps,
    missing_reference_files,
    missing_render_tools,
    validate_workbook_against_spec,
)
from sample_workbook import export_sample_workbook  # noqa: E402

REFERENCES_DIR = REPO_ROOT / "references"

REFERENCE_DIMENSIONS = {
    "Derechos e impuestos": (1853, 91),
    "Nomenclatura": (1531, 124),
    "Restricciones": (1492, 147),
    "Cuotas": (1460, 137),
}


def _tooling_skip_reason() -> str | None:
    tool_gaps = missing_render_tools()
    if tool_gaps:
        return (
            f"Visual comparison requires {', '.join(tool_gaps)} (LibreOffice + poppler-utils), "
            "which are not installed in this environment."
        )
    dep_gaps = missing_python_deps()
    if dep_gaps:
        return (
            f"Visual comparison requires the optional 'visual' extra (missing: {', '.join(dep_gaps)}). "
            "Install with `pip install -e .[visual]`."
        )
    return None


def _design_skip_reason() -> str | None:
    reason = _tooling_skip_reason()
    if reason:
        return reason
    ref_gaps = missing_reference_files(REFERENCES_DIR)
    if ref_gaps:
        return (
            "Visual comparison requires reference screenshots that are not present in this "
            f"checkout: {', '.join(ref_gaps)}. Add them under references/ (see references/README.md) to "
            "enable this test."
        )
    return None


def _baseline_skip_reason() -> str | None:
    reason = _tooling_skip_reason()
    if reason:
        return reason
    if not (BASELINE_DIR / BASELINE_MANIFEST_NAME).is_file():
        return (
            f"No renderer baseline under {BASELINE_DIR}. Generate one explicitly with "
            "`python scripts/update_visual_baseline.py --confirm`."
        )
    return None


@pytest.mark.visual
def test_design_reference_comparison_is_within_tolerant_thresholds(tmp_path):
    """Mode A: tolerant, cross-engine comparison with the Excel screenshots."""
    reason = _design_skip_reason()
    if reason:
        pytest.skip(reason)

    workbook_path = export_sample_workbook(tmp_path / "sample_workbook.xlsx")
    # Complementary exact check: pixels across two engines cannot prove headers,
    # column order or styles, but openpyxl can.
    validate_workbook_against_spec(workbook_path)

    results = compare_workbook_to_references(workbook_path, REFERENCES_DIR, tmp_path / "visual-diagnostics")

    assert set(results.keys()) == set(SHEET_REFERENCE_FILES.keys())
    for sheet_name, metrics in results.items():
        assert metrics["reference_dimensions"] == REFERENCE_DIMENSIONS[sheet_name]

    failures = {
        sheet_name: metrics
        for sheet_name, metrics in results.items()
        if metrics["mean_abs_diff"] > DESIGN_MEAN_DIFF_THRESHOLD or metrics["ssim"] < DESIGN_SSIM_THRESHOLD
    }
    assert not failures, f"Design-reference regression (diagnostics under {tmp_path}): {failures}"


@pytest.mark.visual
def test_renderer_baseline_comparison_is_strict(tmp_path):
    """Mode B: strict, same-engine comparison with the approved baseline."""
    reason = _baseline_skip_reason()
    if reason:
        pytest.skip(reason)

    workbook_path = export_sample_workbook(tmp_path / "sample_workbook.xlsx")
    try:
        results = compare_workbook_to_baseline(workbook_path, BASELINE_DIR, tmp_path / "baseline-diagnostics")
    except BaselineUnavailable as error:  # pragma: no cover - depends on the checkout
        pytest.fail(str(error))

    assert set(results.keys()) == set(SHEET_REFERENCE_FILES.keys())
    failures = {
        sheet_name: metrics
        for sheet_name, metrics in results.items()
        if metrics["mean_abs_diff"] > BASELINE_MEAN_DIFF_THRESHOLD or metrics["ssim"] < BASELINE_SSIM_THRESHOLD
    }
    assert not failures, (
        "Renderer-baseline regression. Inspect the diffs under "
        f"{tmp_path}/baseline-diagnostics and, only if the change is intended, run "
        f"`python scripts/update_visual_baseline.py --confirm`: {failures}"
    )


@pytest.mark.visual
def test_material_regression_fails_the_baseline_comparison(tmp_path):
    """Negative control: a deliberate formatting regression must be detected."""
    reason = _baseline_skip_reason()
    if reason:
        pytest.skip(reason)

    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

    workbook_path = export_sample_workbook(tmp_path / "regressed_workbook.xlsx")
    workbook = load_workbook(workbook_path)
    worksheet = workbook["Derechos e impuestos"]
    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor="C00000")
    workbook.save(workbook_path)

    results = compare_workbook_to_baseline(workbook_path, BASELINE_DIR, tmp_path / "negative-diagnostics")

    metrics = results["Derechos e impuestos"]
    assert (
        metrics["mean_abs_diff"] > BASELINE_MEAN_DIFF_THRESHOLD
        or metrics["ssim"] < BASELINE_SSIM_THRESHOLD
    ), f"A red header row must fail the baseline comparison, but it reported {metrics}"
    # Unrelated sheets must stay clean, so the signal is specific.
    assert results["Cuotas"]["ssim"] >= BASELINE_SSIM_THRESHOLD
    # The approved baseline must never be rewritten by a test run.
    assert (BASELINE_DIR / BASELINE_MANIFEST_NAME).is_file()
