"""Unit tests for SAT tariff scraper."""

import json
import os
import re
import tempfile
import unittest
from unittest import mock

import main as main_module
from bs4 import BeautifulSoup
from openpyxl import load_workbook

from main import build_argument_parser, load_hs_codes_from_file
from sat_scraper import (
    INPUT_INDEX_KEY,
    SECTION_LABELS,
    SECTION_ROWS_KEY,
    SECTION_STATUS_KEY,
    SATTariffScraper,
)


class TestSATTariffScraper(unittest.TestCase):
    def setUp(self):
        self.scraper = SATTariffScraper()

    @staticmethod
    def _result_for(hs_code, status, input_index=None):
        return {
            INPUT_INDEX_KEY: input_index,
            "HS_Code": hs_code,
            "Status": status,
            "Sections": {section_label: {} for section_label in SECTION_LABELS},
            "Section_Statuses": (
                {section_label: "Success" for section_label in SECTION_LABELS}
                if status == "Success"
                else {}
            ),
        }

    @staticmethod
    def _header_rows_for_sheet(sheet):
        coordinate = str(sheet.freeze_panes or "A2")
        match = re.search(r"(\d+)$", coordinate)
        return int(match.group(1)) - 1 if match else 1

    @classmethod
    def _data_start_row_for_sheet(cls, sheet):
        return cls._header_rows_for_sheet(sheet) + 1

    @staticmethod
    def _merged_parent_value(sheet, row_index, column_index):
        value = sheet.cell(row=row_index, column=column_index).value
        if value is not None:
            return value
        for merged_range in sheet.merged_cells.ranges:
            if (
                merged_range.min_row <= row_index <= merged_range.max_row
                and merged_range.min_col <= column_index <= merged_range.max_col
            ):
                return sheet.cell(merged_range.min_row, merged_range.min_col).value
        return None

    @classmethod
    def _visible_headers(cls, sheet):
        if cls._header_rows_for_sheet(sheet) == 1:
            return [sheet.cell(row=1, column=index).value for index in range(1, sheet.max_column + 1)]

        headers = []
        for column_index in range(1, sheet.max_column + 1):
            parent = cls._merged_parent_value(sheet, 1, column_index)
            child = sheet.cell(row=2, column=column_index).value
            if parent == "Unidades de medida" and child:
                headers.append(f"Unidades de medida - {child}")
            else:
                headers.append(child or parent)
        return headers

    def test_load_hs_codes_from_file_supports_arbitrary_counts_without_cap(self):
        for count in (0, 1, 19, 20, 25):
            with self.subTest(count=count):
                codes = [f"{1000000000 + index:010d}" for index in range(count)]
                with tempfile.NamedTemporaryFile("w", delete=False, encoding="utf-8") as temp_file:
                    temp_file.write("number\n")
                    temp_file.write("\n".join(codes))
                    input_path = temp_file.name

                try:
                    self.assertEqual(load_hs_codes_from_file(input_path), codes)
                finally:
                    os.remove(input_path)

    def test_argument_parser_accepts_resume_delay_and_retry_options(self):
        args = build_argument_parser().parse_args(
            [
                "hs_codes.txt",
                "output.xlsx",
                "--resume",
                "--state-file",
                "state.json",
                "--delay-between-codes",
                "1.5",
                "--max-retries",
                "3",
                "--retry-backoff",
                "4",
            ]
        )

        self.assertEqual(args.hs_codes_file, "hs_codes.txt")
        self.assertEqual(args.output_file, "output.xlsx")
        self.assertTrue(args.resume)
        self.assertEqual(args.state_file, "state.json")
        self.assertEqual(args.delay_between_codes, 1.5)
        self.assertEqual(args.max_retries, 3)
        self.assertEqual(args.retry_backoff, 4)

    def test_main_exits_when_no_default_hs_codes_are_configured(self):
        with mock.patch.object(main_module, "HS_CODES", []):
            with self.assertRaises(SystemExit) as raised:
                main_module.main([])

        self.assertEqual(raised.exception.code, 1)

    def test_sections_are_processed_in_requested_order(self):
        self.assertEqual(
            SECTION_LABELS,
            ("Derechos e impuestos", "Nomenclatura", "Restricciones", "Cuotas"),
        )

    def test_parse_nomenclature_extracts_named_fields_and_messages(self):
        soup = BeautifulSoup(
            """
            <table>
              <tr><td>Sección</td><td>Sección I</td></tr>
              <tr><td>Capítulo:</td><td>Capítulo 01</td></tr>
              <tr><td>Fecha inicio de vigencia:</td><td>2024-01-01</td></tr>
              <tr><td>Fecha fin de vigencia:</td><td>2024-12-31</td></tr>
            </table>
            <h3>Código de Mercancías</h3>
            <table>
              <tr><th>Código</th><th>Descripción</th><th>Notas</th></tr>
              <tr><td>0101</td><td>Caballos</td><td>Vivo</td></tr>
            </table>
            <h3>Códigos adicionales</h3>
            <p>No se han encontrado códigos adicionales asociados al inciso consultado</p>
            <h3>Unidades de medida</h3>
            <table>
              <tr><th>Código</th><th>Descripción</th><th>Extra</th></tr>
              <tr><td>KGM</td><td>Kilogramo</td><td>Ignorar</td></tr>
            </table>
            <h3>Clasificadores estadísticos</h3>
            <p>No aplica</p>
            <h3>Descripciones mínimas</h3>
            <p>Sin descripción mínima</p>
            <h3>Criterios de Clasificación</h3>
            <p>Según nota legal</p>
            """,
            "html.parser",
        )

        data = self.scraper._parse_nomenclature(soup)
        rows = data[SECTION_ROWS_KEY]

        mercancías_row = next(row for row in rows if row["Table_Name"] == "Código de Mercancías")
        self.assertEqual(mercancías_row["Sección"], "Sección I")
        self.assertEqual(mercancías_row["Capítulo:"], "Capítulo 01")
        self.assertEqual(mercancías_row["Fecha inicio de vigencia:"], "2024-01-01")
        self.assertEqual(mercancías_row["Fecha fin de vigencia:"], "2024-12-31")
        self.assertEqual(mercancías_row["Código"], "0101")
        self.assertEqual(mercancías_row["Descripción"], "Caballos")
        self.assertEqual(mercancías_row["Notas"], "Vivo")

        additional_codes_row = next(row for row in rows if row["Table_Name"] == "Códigos adicionales")
        self.assertEqual(
            additional_codes_row["Message"],
            "No se han encontrado códigos adicionales asociados al inciso consultado",
        )

        units_row = next(row for row in rows if row["Table_Name"] == "Unidades de medida")
        self.assertEqual(units_row["Código"], "KGM")
        self.assertEqual(units_row["Descripción"], "Kilogramo")
        self.assertNotIn("Extra", units_row)

        self.assertIn(
            {
                "Sección": "Sección I",
                "Capítulo:": "Capítulo 01",
                "Fecha inicio de vigencia:": "2024-01-01",
                "Fecha fin de vigencia:": "2024-12-31",
                "Record_Type": "Clasificadores estadísticos",
                "Table_Name": "Clasificadores estadísticos",
                "Content": "No aplica",
            },
            rows,
        )
        self.assertIsNone(data[SECTION_STATUS_KEY])

    def test_parse_duties_preserves_duplicate_and_agreement_rows(self):
        soup = BeautifulSoup(
            """
            <h3>TRATAMIENTO GENERAL</h3>
            <table>
              <tr>
                <th>Código</th><th>Descripción</th><th>Código adicional</th>
                <th>Valor</th><th>Código de cuota</th>
              </tr>
              <tr><td>A1</td><td>IVA</td><td>AD1</td><td>12%</td><td>CQ1</td></tr>
              <tr><td>A1</td><td>IVA</td><td>AD1</td><td>12%</td><td>CQ1</td></tr>
            </table>
            <h3>TLC Guatemala</h3>
            <table>
              <tr>
                <th>Código</th><th>Descripción</th><th>Código adicional</th>
                <th>Valor</th><th>Código de cuota</th>
              </tr>
              <tr><td>B2</td><td>Preferencial</td><td>AD2</td><td>0%</td><td>CQ2</td></tr>
            </table>
            """,
            "html.parser",
        )

        data = self.scraper._parse_standard_section(soup)
        rows = data[SECTION_ROWS_KEY]

        self.assertEqual(len(rows), 3)
        self.assertEqual(sum(1 for row in rows if row["Código"] == "A1"), 2)
        self.assertEqual(rows[0]["Table_Name"], "TRATAMIENTO GENERAL")
        self.assertEqual(rows[0]["Código de cuota"], "CQ1")
        self.assertEqual(rows[2]["Table_Name"], "TLC Guatemala")
        self.assertEqual(rows[2]["Valor"], "0%")

    def test_parse_restrictions_general_rows_include_quota_code(self):
        soup = BeautifulSoup(
            """
            <h3>TRATAMIENTO GENERAL</h3>
            <table>
              <tr>
                <th>Código</th><th>Descripción</th><th>Código adicional</th>
                <th>Valor</th><th>Código de cuota</th>
              </tr>
              <tr><td>R1</td><td>Licencia previa</td><td>AD3</td><td>Obligatoria</td><td>CQR</td></tr>
            </table>
            """,
            "html.parser",
        )

        data = self.scraper._parse_standard_section(soup)
        self.assertEqual(
            data[SECTION_ROWS_KEY],
            [
                {
                    "Table_Name": "TRATAMIENTO GENERAL",
                    "Código": "R1",
                    "Descripción": "Licencia previa",
                    "Código adicional": "AD3",
                    "Valor": "Obligatoria",
                    "Código de cuota": "CQR",
                }
            ],
        )

    def test_parse_quotas_empty_message_when_no_rows_exist(self):
        soup = BeautifulSoup(
            """
            <div>No se han encontrado cuotas/contingentes para el inciso consultado</div>
            """,
            "html.parser",
        )

        data = self.scraper._parse_standard_section(
            soup,
            empty_message="No se han encontrado cuotas/contingentes para el inciso consultado",
        )

        self.assertEqual(
            data[SECTION_ROWS_KEY],
            [
                {
                    "Table_Name": "TRATAMIENTO GENERAL",
                    "Message": "No se han encontrado cuotas/contingentes para el inciso consultado",
                }
            ],
        )
        self.assertIsNone(data[SECTION_STATUS_KEY])

    def test_format_quota_message_does_not_duplicate_prefix(self):
        prefixed = (
            "Resultados de la búsqueda: "
            "No se han encontrado cuotas/contingentes para el inciso consultado"
        )
        self.assertEqual(self.scraper._format_quota_message(prefixed), prefixed)

    def test_scrape_hs_code_reflects_non_success_section_status(self):
        expected_data = {
            "Derechos e impuestos": {SECTION_ROWS_KEY: [{"Table_Name": "TRATAMIENTO GENERAL"}]},
            "Nomenclatura": {SECTION_ROWS_KEY: [], SECTION_STATUS_KEY: "No data found"},
            "Restricciones": {SECTION_ROWS_KEY: [{"Table_Name": "TRATAMIENTO GENERAL"}]},
            "Cuotas": {SECTION_ROWS_KEY: [{"Table_Name": "TRATAMIENTO GENERAL", "Message": "Sin cuota"}]},
        }

        self.scraper.check_for_captcha = lambda: False
        self.scraper.enter_hs_code = lambda hs_code: True
        self.scraper.click_search_button = lambda: True
        self.scraper.extract_section_in_new_tab = lambda section_label: expected_data[section_label]

        result = self.scraper.scrape_hs_code("0101210000")

        self.assertEqual(result["Section_Statuses"]["Nomenclatura"], "No data found")
        self.assertEqual(result["Status"], "Section issues: Nomenclatura: No data found")

    def test_scrape_multiple_continues_after_failed_code_with_bounded_retries(self):
        scraper = SATTariffScraper(delay_between_codes=0, max_retries=2, retry_backoff=0)
        scraper.navigate_to_portal = lambda: None
        scraper.export_to_excel = lambda output_file: None
        attempts = []
        outcomes = {
            "0101210000": [
                self._result_for("0101210000", "Failed to submit HS code"),
                self._result_for("0101210000", "Failed to submit HS code"),
                self._result_for("0101210000", "Failed to submit HS code"),
            ],
            "0102210000": [self._result_for("0102210000", "Success")],
        }

        def fake_scrape_hs_code(hs_code, input_index=None):
            attempts.append(hs_code)
            result = outcomes[hs_code].pop(0)
            result[INPUT_INDEX_KEY] = input_index
            return result

        scraper.scrape_hs_code = fake_scrape_hs_code

        scraper.scrape_multiple(
            ["0101210000", "0102210000"],
            output_file="/tmp/test-output.xlsx",
        )

        self.assertEqual(
            attempts,
            ["0101210000", "0101210000", "0101210000", "0102210000"],
        )
        self.assertEqual([result["HS_Code"] for result in scraper.results], ["0101210000", "0102210000"])
        self.assertEqual(scraper.results[0]["Status"], "Failed to submit HS code")
        self.assertEqual(scraper.results[1]["Status"], "Success")

    def test_scrape_multiple_does_not_retry_captcha_failures(self):
        scraper = SATTariffScraper(delay_between_codes=0, max_retries=3, retry_backoff=0)
        scraper.navigate_to_portal = lambda: None
        scraper.export_to_excel = lambda output_file: None
        attempts = []

        def fake_scrape_hs_code(hs_code, input_index=None):
            attempts.append(hs_code)
            return self._result_for(hs_code, "CAPTCHA solving failed", input_index=input_index)

        scraper.scrape_hs_code = fake_scrape_hs_code

        scraper.scrape_multiple(
            ["0101210000", "0102210000"],
            output_file="/tmp/test-output.xlsx",
        )

        self.assertEqual(attempts, ["0101210000"])
        self.assertEqual(len(scraper.results), 1)
        self.assertEqual(scraper.results[0]["Status"], "CAPTCHA solving failed")

    def test_export_to_excel_creates_reference_layout_workbook(self):
        self.scraper.results = [
            {
                "HS_Code": "0101210000",
                "Status": "Success",
                "Sections": {
                    "Derechos e impuestos": {
                        SECTION_ROWS_KEY: [
                            {
                                "Table_Name": "TRATAMIENTO GENERAL",
                                "Código": "DAI",
                                "Descripción": "Derecho arancelario",
                                "Código adicional": "AD1",
                                "Valor": "5%",
                                "Código de cuota": "CQ1",
                            },
                            {
                                "Table_Name": "TRATAMIENTO GENERAL",
                                "Código": "IVA",
                                "Descripción": "Impuesto al valor agregado",
                                "Código adicional": "AD1",
                                "Valor": "12%",
                                "Código de cuota": "CQ1",
                            },
                            {
                                "Table_Name": "TLC México",
                                "Código": "DAI",
                                "Descripción": "Preferencial",
                                "Código adicional": "AD1",
                                "Valor": "0%",
                                "Código de cuota": "CQ1",
                            },
                            {
                                "Table_Name": "TLC Chile",
                                "Código": "DAI",
                                "Descripción": "Preferencial",
                                "Código adicional": "AD1",
                                "Valor": "1%",
                                "Código de cuota": "CQ1",
                            },
                        ]
                    },
                    "Nomenclatura": {
                        SECTION_ROWS_KEY: [
                            {
                                "Table_Name": "Código de Mercancías",
                                "Record_Type": "Código de Mercancías",
                                "Sección": "Sección I",
                                "Capítulo:": "Capítulo 01",
                                "Fecha inicio de vigencia:": "2024-01-01",
                                "Fecha fin de vigencia:": "2024-12-31",
                                "Código": "0101",
                                "Descripción": "Caballos",
                                "Notas": "Vivo",
                            },
                            {
                                "Table_Name": "Códigos adicionales",
                                "Record_Type": "Códigos adicionales",
                                "Sección": "Sección I",
                                "Capítulo:": "Capítulo 01",
                                "Fecha inicio de vigencia:": "2024-01-01",
                                "Fecha fin de vigencia:": "2024-12-31",
                                "Message": "No se han encontrado códigos adicionales asociados al inciso consultado",
                            },
                            {
                                "Table_Name": "Unidades de medida",
                                "Record_Type": "Unidades de medida",
                                "Sección": "Sección I",
                                "Capítulo:": "Capítulo 01",
                                "Fecha inicio de vigencia:": "2024-01-01",
                                "Fecha fin de vigencia:": "2024-12-31",
                                "Código": "KGM",
                                "Descripción": "Kilogramo",
                            }
                        ]
                    },
                    "Restricciones": {
                        SECTION_ROWS_KEY: [
                            {
                                "Table_Name": "TRATAMIENTO GENERAL",
                                "Código": "R1",
                                "Descripción": "Licencia previa",
                                "Código adicional": "AD3",
                                "Valor": "Obligatoria",
                                "Código de cuota": "CQR",
                            }
                            ,
                            {
                                "Table_Name": "TRATAMIENTO GENERAL",
                                "Código": "R1",
                                "Descripción": "Licencia previa",
                                "Código adicional": "AD3",
                                "Valor": "Obligatoria",
                                "Código de cuota": "CQR",
                            },
                        ]
                    },
                    "Cuotas": {
                        SECTION_ROWS_KEY: [
                            {
                                "Table_Name": "TRATAMIENTO GENERAL",
                                "Message": "No se han encontrado cuotas/contingentes para el inciso consultado",
                            }
                        ]
                    },
                },
                "Section_Statuses": {section_label: "Success" for section_label in SECTION_LABELS},
            }
        ]

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as temp_file:
            output_file = temp_file.name

        try:
            self.scraper.export_to_excel(output_file)
            workbook = load_workbook(output_file)

            self.assertEqual(workbook.sheetnames, list(SECTION_LABELS))

            duties_sheet = workbook["Derechos e impuestos"]
            self.assertEqual(
                self._visible_headers(duties_sheet)[:9],
                [
                    "HS_Code",
                    "Status",
                    "Overall_Status",
                    "DAI_GENERAL",
                    "IVA_GENERAL",
                    "DAI_MX",
                    "DAI_CL",
                    "Código adicional",
                    "Código de cuota",
                ],
            )
            self.assertIn("D1:E1", {str(cell_range) for cell_range in duties_sheet.merged_cells.ranges})
            self.assertEqual(duties_sheet["D1"].value, "GENERAL")
            self.assertEqual(duties_sheet["F1"].value, "MX")
            self.assertEqual(duties_sheet["G1"].value, "CL")
            self.assertEqual(duties_sheet["A3"].value, "0101210000")
            self.assertEqual(duties_sheet["D3"].value, "5%")
            self.assertEqual(duties_sheet["E3"].value, "12%")
            self.assertEqual(duties_sheet["F3"].value, "0%")
            self.assertEqual(duties_sheet["G3"].value, "1%")
            self.assertEqual(duties_sheet["H3"].value, "AD1")
            self.assertEqual(duties_sheet["I3"].value, "CQ1")

            nomenclature_sheet = workbook["Nomenclatura"]
            self.assertEqual(
                self._visible_headers(nomenclature_sheet)[:10],
                [
                    "HS_Code",
                    "Status",
                    "Overall_Status",
                    "Sección",
                    "Capítulo:",
                    "Fecha inicio de vigencia:",
                    "Fecha fin de vigencia:",
                    "Códigos adicionales",
                    "Unidades de medida - Código",
                    "Unidades de medida - Descripción",
                ],
            )
            self.assertIn("I1:J1", {str(cell_range) for cell_range in nomenclature_sheet.merged_cells.ranges})
            self.assertEqual(nomenclature_sheet["I1"].value, "Unidades de medida")
            self.assertEqual(nomenclature_sheet["I2"].value, "Código")
            self.assertEqual(nomenclature_sheet["J2"].value, "Descripción")
            self.assertEqual(nomenclature_sheet["A3"].value, "0101210000")
            self.assertEqual(nomenclature_sheet["H4"].value, "No se han encontrado códigos adicionales asociados al inciso consultado")
            self.assertEqual(nomenclature_sheet["I5"].value, "KGM")
            self.assertEqual(nomenclature_sheet["J5"].value, "Kilogramo")

            restrictions_sheet = workbook["Restricciones"]
            self.assertEqual(
                self._visible_headers(restrictions_sheet)[:8],
                [
                    "HS_Code",
                    "Status",
                    "Overall_Status",
                    "Código",
                    "Descripción",
                    "Código adicional",
                    "Valor",
                    "Código de cuota",
                ],
            )
            self.assertEqual(restrictions_sheet["A2"].value, "0101210000")
            self.assertEqual(restrictions_sheet["H2"].value, "CQR")
            self.assertEqual(restrictions_sheet["A3"].value, "0101210000")
            self.assertEqual(restrictions_sheet["H3"].value, "CQR")

            quotas_sheet = workbook["Cuotas"]
            self.assertEqual(quotas_sheet["A2"].value, "0101210000")
            self.assertEqual(quotas_sheet["B2"].value, "Success")
            self.assertEqual(self._visible_headers(quotas_sheet)[:4], ["HS_Code", "Status", "Overall_Status", "Resultado"])
            self.assertEqual(
                quotas_sheet["D2"].value,
                "Resultados de la búsqueda: No se han encontrado cuotas/contingentes para el inciso consultado",
            )
            self.assertEqual(duties_sheet.freeze_panes, "A3")
            self.assertEqual(nomenclature_sheet.freeze_panes, "A3")
            self.assertEqual(restrictions_sheet.freeze_panes, "A2")
            self.assertEqual(quotas_sheet.freeze_panes, "A2")
            self.assertEqual(duties_sheet["A1"].fill.fgColor.rgb[-6:], "4472C4")
            self.assertTrue(duties_sheet["A1"].font.bold)
            self.assertEqual(duties_sheet["A1"].font.color.rgb[-6:], "FFFFFF")
        finally:
            if os.path.exists(output_file):
                os.remove(output_file)

    def test_export_to_excel_uses_overall_status_for_unattempted_sections(self):
        self.scraper.results = [
            {
                "HS_Code": "0101210000",
                "Status": "Failed to submit HS code",
                "Sections": {section_label: {} for section_label in SECTION_LABELS},
                "Section_Statuses": {},
            }
        ]

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as temp_file:
            output_file = temp_file.name

        try:
            self.scraper.export_to_excel(output_file)
            workbook = load_workbook(output_file)

            for section_label in SECTION_LABELS:
                sheet = workbook[section_label]
                data_row = self._data_start_row_for_sheet(sheet)
                self.assertEqual(sheet.cell(row=data_row, column=1).value, "0101210000")
                self.assertEqual(sheet.cell(row=data_row, column=2).value, "Failed to submit HS code")
                self.assertEqual(sheet.cell(row=data_row, column=3).value, "Failed to submit HS code")
        finally:
            if os.path.exists(output_file):
                os.remove(output_file)

    def test_run_resume_skips_successful_codes_retries_failed_codes_and_preserves_order(self):
        successful = self._result_for("0101210000", "Success", input_index=0)
        failed = self._result_for("0102210000", "Failed to submit HS code", input_index=1)

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as output_handle:
            output_file = output_handle.name
        with tempfile.NamedTemporaryFile(suffix=".json", delete=False) as state_handle:
            state_file = state_handle.name

        try:
            seed_scraper = SATTariffScraper()
            seed_scraper.input_order = ["0101210000", "0102210000", "0103210000"]
            seed_scraper.results = [successful, failed]
            seed_scraper.save_state(state_file, output_file)

            scraper = SATTariffScraper(delay_between_codes=0, max_retries=0, retry_backoff=0)
            scraper.start_browser = lambda: None
            scraper.close_browser = lambda: None
            scraper.navigate_to_portal = lambda: None
            attempted_codes = []

            def fake_scrape_hs_code(hs_code, input_index=None):
                attempted_codes.append(hs_code)
                return self._result_for(hs_code, "Success", input_index=input_index)

            scraper.scrape_hs_code = fake_scrape_hs_code

            scraper.run(
                ["0101210000", "0102210000", "0103210000"],
                output_file=output_file,
                resume=True,
                state_file=state_file,
            )

            self.assertEqual(attempted_codes, ["0102210000", "0103210000"])
            self.assertEqual(
                [result["HS_Code"] for result in scraper.results],
                ["0101210000", "0102210000", "0103210000"],
            )
            self.assertEqual(
                [result["Status"] for result in scraper.results],
                ["Success", "Success", "Success"],
            )

            workbook = load_workbook(output_file)
            self.assertEqual(workbook.sheetnames, list(SECTION_LABELS))
            for section_label in SECTION_LABELS:
                sheet = workbook[section_label]
                start_row = self._data_start_row_for_sheet(sheet)
                self.assertEqual(
                    [sheet.cell(row=row, column=1).value for row in range(start_row, start_row + 3)],
                    ["0101210000", "0102210000", "0103210000"],
                )
        finally:
            if os.path.exists(output_file):
                os.remove(output_file)
            if os.path.exists(state_file):
                os.remove(state_file)

    def test_run_resume_tracks_duplicate_hs_codes_by_input_position(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as output_handle:
            output_file = output_handle.name
        with tempfile.NamedTemporaryFile(suffix=".json", delete=False) as state_handle:
            state_file = state_handle.name

        try:
            legacy_state = {
                "output_file": output_file,
                "input_order": ["0101210000", "0101210000"],
                "results": [
                    {
                        "HS_Code": "0101210000",
                        "Status": "Success",
                        "Sections": {section_label: {} for section_label in SECTION_LABELS},
                        "Section_Statuses": {
                            section_label: "Success" for section_label in SECTION_LABELS
                        },
                    }
                ],
            }
            with open(state_file, "w", encoding="utf-8") as file_handle:
                json.dump(legacy_state, file_handle)

            scraper = SATTariffScraper(delay_between_codes=0, max_retries=0, retry_backoff=0)
            scraper.start_browser = lambda: None
            scraper.close_browser = lambda: None
            scraper.navigate_to_portal = lambda: None
            attempted_codes = []

            def fake_scrape_hs_code(hs_code, input_index=None):
                attempted_codes.append((hs_code, input_index))
                return self._result_for(hs_code, "Success", input_index=input_index)

            scraper.scrape_hs_code = fake_scrape_hs_code

            scraper.run(
                ["0101210000", "0101210000"],
                output_file=output_file,
                resume=True,
                state_file=state_file,
            )

            self.assertEqual(attempted_codes, [("0101210000", 1)])
            self.assertEqual(
                [result[INPUT_INDEX_KEY] for result in scraper.results],
                [0, 1],
            )
            self.assertEqual(
                [result["HS_Code"] for result in scraper.results],
                ["0101210000", "0101210000"],
            )
        finally:
            if os.path.exists(output_file):
                os.remove(output_file)
            if os.path.exists(state_file):
                os.remove(state_file)

    def test_run_resume_ignores_malformed_state_file(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as output_handle:
            output_file = output_handle.name
        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as state_handle:
            state_handle.write("{not valid json")
            state_file = state_handle.name

        try:
            scraper = SATTariffScraper(delay_between_codes=0, max_retries=0, retry_backoff=0)
            scraper.start_browser = lambda: None
            scraper.close_browser = lambda: None
            scraper.navigate_to_portal = lambda: None
            attempted_codes = []

            def fake_scrape_hs_code(hs_code, input_index=None):
                attempted_codes.append((hs_code, input_index))
                return self._result_for(hs_code, "Success", input_index=input_index)

            scraper.scrape_hs_code = fake_scrape_hs_code

            scraper.run(
                ["0101210000", "0102210000"],
                output_file=output_file,
                resume=True,
                state_file=state_file,
            )

            self.assertEqual(
                attempted_codes,
                [("0101210000", 0), ("0102210000", 1)],
            )
            self.assertEqual(len(scraper.results), 2)
            self.assertEqual(scraper.input_order, ["0101210000", "0102210000"])
        finally:
            if os.path.exists(output_file):
                os.remove(output_file)
            if os.path.exists(state_file):
                os.remove(state_file)

    def test_load_state_failure_clears_previous_results_and_order(self):
        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as state_handle:
            state_handle.write("{broken json")
            state_file = state_handle.name

        try:
            scraper = SATTariffScraper()
            scraper.input_order = ["stale"]
            scraper.results = [self._result_for("0101210000", "Success", input_index=0)]

            self.assertFalse(scraper.load_state(state_file))
            self.assertEqual(scraper.input_order, [])
            self.assertEqual(scraper.results, [])
        finally:
            if os.path.exists(state_file):
                os.remove(state_file)

    def test_load_state_rejects_non_mapping_payloads(self):
        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as state_handle:
            json.dump([], state_handle)
            state_file = state_handle.name

        try:
            scraper = SATTariffScraper()
            self.assertFalse(scraper.load_state(state_file))
            self.assertEqual(scraper.input_order, [])
            self.assertEqual(scraper.results, [])
        finally:
            if os.path.exists(state_file):
                os.remove(state_file)

    def test_run_resume_without_explicit_hs_codes_uses_state_input_order(self):
        state_payload = {
            "output_file": "ignored.xlsx",
            "input_order": ["0101210000", "0102210000"],
            "results": [self._result_for("0101210000", "Success", input_index=0)],
        }

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as output_handle:
            output_file = output_handle.name
        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as state_handle:
            json.dump(state_payload, state_handle)
            state_file = state_handle.name

        try:
            scraper = SATTariffScraper(delay_between_codes=0, max_retries=0, retry_backoff=0)
            scraper.start_browser = lambda: None
            scraper.close_browser = lambda: None
            scraper.navigate_to_portal = lambda: None
            attempted_codes = []

            def fake_scrape_hs_code(hs_code, input_index=None):
                attempted_codes.append((hs_code, input_index))
                return self._result_for(hs_code, "Success", input_index=input_index)

            scraper.scrape_hs_code = fake_scrape_hs_code

            scraper.run(
                [],
                output_file=output_file,
                resume=True,
                state_file=state_file,
                use_state_input_order=True,
            )

            self.assertEqual(attempted_codes, [("0102210000", 1)])
            self.assertEqual(
                [result["HS_Code"] for result in scraper.results],
                ["0101210000", "0102210000"],
            )
        finally:
            if os.path.exists(output_file):
                os.remove(output_file)
            if os.path.exists(state_file):
                os.remove(state_file)

    def test_run_with_no_codes_creates_empty_four_sheet_workbook(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as output_handle:
            output_file = output_handle.name
        with tempfile.NamedTemporaryFile(suffix=".json", delete=False) as state_handle:
            state_file = state_handle.name

        try:
            scraper = SATTariffScraper()
            scraper.start_browser = mock.Mock()
            scraper.run([], output_file=output_file, state_file=state_file)

            scraper.start_browser.assert_not_called()
            workbook = load_workbook(output_file)
            self.assertEqual(workbook.sheetnames, list(SECTION_LABELS))
        finally:
            if os.path.exists(output_file):
                os.remove(output_file)
            if os.path.exists(state_file):
                os.remove(state_file)


if __name__ == "__main__":
    unittest.main()
