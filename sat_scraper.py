"""SAT tariff scraper with manual CAPTCHA handling and section extraction."""

import json
import logging
import re
import time
from pathlib import Path
from typing import Dict, List, Optional

import pandas as pd
from bs4 import BeautifulSoup, NavigableString, Tag
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from selenium import webdriver
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait
from webdriver_manager.chrome import ChromeDriverManager

from config import (
    MAX_RETRIES,
    PAGE_LOAD_DELAY,
    REQUEST_DELAY,
    RETRY_BACKOFF,
    SAT_BASE_URL,
    WAIT_TIMEOUT,
)

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)

PAGE_RELOAD_DELAY = PAGE_LOAD_DELAY
CAPTCHA_TIMEOUT = 300
SECTION_DELAY = REQUEST_DELAY
DUTIES_TAB_DELAY = 5

SECTION_LABELS = (
    "Derechos e impuestos",
    "Nomenclatura",
    "Restricciones",
    "Cuotas",
)
INPUT_INDEX_KEY = "Input_Index"
SECTION_STATUS_KEY = "__section_status__"
SECTION_ROWS_KEY = "__section_rows__"


class SATTariffScraper:
    """Scrape every requested SAT tariff section for each HS code."""

    STANDARD_SECTION_COLUMNS = (
        "Código",
        "Descripción",
        "Código adicional",
        "Valor",
        "Código de cuota",
    )
    NOMENCLATURE_BLOCK_LABELS = (
        "Código de Mercancías",
        "Códigos adicionales",
        "Unidades de medida",
        "Clasificadores estadísticos",
        "Descripciones mínimas",
        "Criterios de Clasificación",
    )

    def __init__(
        self,
        headless=False,
        manual_captcha=True,
        captcha_timeout=CAPTCHA_TIMEOUT,
        delay_between_codes=REQUEST_DELAY,
        max_retries=MAX_RETRIES,
        retry_backoff=RETRY_BACKOFF,
    ):
        self.base_url = SAT_BASE_URL
        self.consulta_url = (
            "https://farm2.sat.gob.gt/saqbe-arancel-publico"
            "/aduana/arancel/consulta/consulta.jsf"
        )
        self.driver = None
        self.wait = None
        self.results: List[Dict] = []
        self.input_order: List[str] = []
        self.headless = headless
        self.manual_captcha = manual_captcha
        self.captcha_timeout = captcha_timeout
        self.delay_between_codes = max(0, delay_between_codes)
        self.max_retries = max(0, max_retries)
        self.retry_backoff = max(0, retry_backoff)

    def start_browser(self):
        options = webdriver.ChromeOptions()
        if self.headless:
            options.add_argument("--headless")
        options.add_argument("--no-sandbox")
        options.add_argument("--disable-dev-shm-usage")
        options.add_argument("--disable-popup-blocking")
        options.add_argument("--disable-blink-features=AutomationControlled")
        options.add_experimental_option("excludeSwitches", ["enable-automation"])
        options.add_experimental_option("useAutomationExtension", False)
        service = Service(ChromeDriverManager().install())
        self.driver = webdriver.Chrome(service=service, options=options)
        self.wait = WebDriverWait(self.driver, WAIT_TIMEOUT)
        logger.info("Browser started successfully")

    def navigate_to_portal(self):
        self.driver.get(self.consulta_url)
        time.sleep(PAGE_RELOAD_DELAY)
        self.wait.until(EC.presence_of_element_located((By.TAG_NAME, "body")))
        logger.info("SAT consultation page loaded")

    def find_captcha_field(self):
        selectors = (
            "//input[@id='frmBuscar:txtKaptcha']",
            "//input[contains(@id, 'Kaptcha')]",
            "//input[contains(translate(@id, 'CAPTCHA', 'captcha'), 'captcha')]",
            "//input[contains(translate(@name, 'CAPTCHA', 'captcha'), 'captcha')]",
        )
        for selector in selectors:
            if self.driver.find_elements(By.XPATH, selector):
                return self.driver.find_elements(By.XPATH, selector)[0]
        return None

    def check_for_captcha(self):
        return self.find_captcha_field() is not None

    def handle_captcha_manual(self):
        logger.warning("CAPTCHA detected. Solve it in the visible browser window.")
        logger.warning("The scraper will continue when the CAPTCHA field disappears.")
        deadline = time.time() + self.captcha_timeout
        while time.time() < deadline:
            if not self.check_for_captcha():
                logger.info("CAPTCHA cleared")
                time.sleep(2)
                return True
            time.sleep(1)
        logger.error("CAPTCHA solving timeout")
        return False

    def _find_first(self, selectors, clickable=False):
        def probe(driver):
            for selector in selectors:
                for element in driver.find_elements(By.XPATH, selector):
                    if clickable and not (element.is_displayed() and element.is_enabled()):
                        continue
                    return element
            return False

        try:
            return self.wait.until(probe)
        except TimeoutException:
            return None

    def find_input_field(self):
        return self._find_first((
            "//input[@id='frmBuscar:txtCodigo']",
            "//input[contains(@id, 'txtCodigoArancelario')]",
            "//input[contains(@id, 'txtCodigo')]",
            "//input[contains(@placeholder, 'arancel')]",
            "//input[contains(@class, 'form-control')]",
        ))

    def enter_hs_code(self, hs_code: str):
        field = self.find_input_field()
        if field is None:
            logger.error("HS Code input field not found")
            return False
        field.clear()
        field.send_keys(hs_code)
        return True

    def click_search_button(self):
        button = self._find_first((
            "//input[@id='frmBuscar:pen']",
            "//input[@type='submit' and @value='Buscar']",
            "//input[@type='submit' and contains(@value, 'Consultar')]",
            "//button[contains(@id, 'btnBuscar')]",
            "//button[contains(text(), 'Consultar')]",
            "//button[contains(text(), 'Buscar')]",
            "//button[@type='submit']",
        ), clickable=True)
        if button is None:
            logger.error("Search button not found")
            return False
        button.click()
        time.sleep(3)
        return True

    @staticmethod
    def _direct_rows(table):
        return [row for row in table.find_all("tr") if row.find_parent("table") is table]

    @staticmethod
    def _direct_cells(row):
        return [cell for cell in row.find_all(["td", "th"]) if cell.find_parent("tr") is row]

    @staticmethod
    def _cell_text(cell):
        return cell.get_text(" ", strip=True)

    @staticmethod
    def _normalize_text(text: str) -> str:
        return re.sub(r"\s+", " ", (text or "")).strip()

    @classmethod
    def _normalize_label(cls, text: str) -> str:
        return cls._normalize_text(text).rstrip(":").casefold()

    def _table_to_rows(self, table, selected_headers=None):
        rows = self._direct_rows(table)
        if not rows:
            return []
        header_cells = self._direct_cells(rows[0])
        headers = [
            self._cell_text(cell) or f"Column_{index}"
            for index, cell in enumerate(header_cells, start=1)
        ]
        extracted_rows = []
        for row in rows[1:]:
            cells = [self._cell_text(cell) for cell in self._direct_cells(row)]
            if not any(cells):
                continue
            row_data = {}
            for index, header in enumerate(headers):
                if selected_headers and header not in selected_headers:
                    continue
                row_data[header] = cells[index] if index < len(cells) else ""
            if row_data:
                extracted_rows.append(row_data)
        return extracted_rows

    def _find_anchor(self, soup, label: str):
        normalized_label = self._normalize_label(label)
        for element in soup.find_all(
            ["h1", "h2", "h3", "h4", "h5", "h6", "strong", "b", "label", "legend", "span", "div", "p", "td", "th"]
        ):
            text = self._normalize_text(element.get_text(" ", strip=True))
            if self._normalize_label(text) == normalized_label:
                return element
        return None

    def _extract_labeled_value(self, soup, label: str) -> str:
        normalized_label = self._normalize_label(label)

        for row in soup.find_all("tr"):
            cells = self._direct_cells(row)
            if len(cells) < 2:
                continue
            if self._normalize_label(self._cell_text(cells[0])) == normalized_label:
                return self._cell_text(cells[1])

        label_pattern = re.compile(
            rf"^{re.escape(self._normalize_text(label).rstrip(':'))}\s*:?\s*(.+)$",
            re.IGNORECASE,
        )
        for element in soup.find_all(["div", "span", "p", "li", "td", "th", "strong", "b"]):
            text = self._normalize_text(element.get_text(" ", strip=True))
            if not text:
                continue
            match = label_pattern.match(text)
            if match:
                return match.group(1).strip()
            if self._normalize_label(text) != normalized_label:
                continue
            for sibling in element.next_siblings:
                if isinstance(sibling, NavigableString):
                    value = self._normalize_text(str(sibling))
                else:
                    value = self._normalize_text(sibling.get_text(" ", strip=True))
                if value:
                    return value
        return ""

    def _find_table_after_label(self, soup, label: str, stop_labels):
        anchor = self._find_anchor(soup, label)
        stop_set = {self._normalize_label(item) for item in stop_labels if item != label}
        if anchor is not None:
            for element in anchor.find_all_next():
                if isinstance(element, Tag):
                    if self._normalize_label(element.get_text(" ", strip=True)) in stop_set:
                        break
                    if element.name == "table":
                        return element
        normalized_label = self._normalize_label(label)
        for table in soup.find_all("table"):
            if self._normalize_label(self._infer_table_name(table, "")) == normalized_label:
                return table
        return None

    def _extract_text_after_label(self, soup, label: str, stop_labels) -> str:
        anchor = self._find_anchor(soup, label)
        if anchor is None:
            return ""
        stop_set = {self._normalize_label(item) for item in stop_labels if item != label}
        for element in anchor.find_all_next():
            if isinstance(element, Tag):
                text = self._normalize_text(element.get_text(" ", strip=True))
                if text and self._normalize_label(text) in stop_set:
                    break
                if element.name == "table":
                    return ""
                if element.find_parent("table") is not None:
                    continue
                if text and self._normalize_label(text) != self._normalize_label(label):
                    return text
            elif isinstance(element, NavigableString):
                text = self._normalize_text(str(element))
                if text:
                    return text
        return ""

    def _infer_table_name(self, table, fallback: str) -> str:
        caption = table.find("caption")
        if caption:
            caption_text = self._normalize_text(caption.get_text(" ", strip=True))
            if caption_text:
                return caption_text

        ignored_headers = {self._normalize_label(column) for column in self.STANDARD_SECTION_COLUMNS}
        for element in table.find_all_previous(
            ["h1", "h2", "h3", "h4", "h5", "h6", "strong", "b", "label", "legend", "div", "span", "p"],
            limit=20,
        ):
            if element.find_parent("table") is not None:
                continue
            text = self._normalize_text(element.get_text(" ", strip=True))
            if not text or len(text) > 120:
                continue
            if self._normalize_label(text) in ignored_headers:
                continue
            return text
        return fallback

    def _extract_matching_message(self, soup, expected_message: str) -> str:
        normalized_expected = self._normalize_label(expected_message)
        for element in soup.find_all(["div", "span", "p", "li", "td", "th"]):
            if element.find_parent("table") is not None:
                continue
            text = self._normalize_text(element.get_text(" ", strip=True))
            if not text:
                continue
            normalized_text = self._normalize_label(text)
            if normalized_expected in normalized_text or normalized_text in normalized_expected:
                return text
        return expected_message

    def _parse_named_table_rows(self, soup, table_name: str, base_fields: Dict, selected_headers=None):
        table = self._find_table_after_label(soup, table_name, self.NOMENCLATURE_BLOCK_LABELS)
        if table is None:
            return []
        rows = []
        for row_data in self._table_to_rows(table, selected_headers=selected_headers):
            rows.append(
                {
                    **base_fields,
                    "Record_Type": table_name,
                    "Table_Name": table_name,
                    **row_data,
                }
            )
        return rows

    def _standard_table_rows(self, soup, empty_message=None):
        rows = []
        normalized_columns = {
            self._normalize_label(column): column for column in self.STANDARD_SECTION_COLUMNS
        }
        for index, table in enumerate(soup.find_all("table"), start=1):
            table_rows = self._direct_rows(table)
            if not table_rows:
                continue
            headers = [self._cell_text(cell) for cell in self._direct_cells(table_rows[0])]
            present_headers = {
                self._normalize_label(header): header for header in headers if header
            }
            if not {"código", "descripción"}.issubset(present_headers):
                continue
            if len(set(present_headers).intersection(normalized_columns)) < 3:
                continue
            table_name = self._infer_table_name(table, f"Tabla {index}")
            for row_data in self._table_to_rows(table):
                normalized_row = {
                    normalized_columns[self._normalize_label(column)]: value
                    for column, value in row_data.items()
                    if self._normalize_label(column) in normalized_columns
                }
                rows.append(
                    {
                        "Table_Name": table_name,
                        **{column: normalized_row.get(column, "") for column in self.STANDARD_SECTION_COLUMNS},
                    }
                )
        if rows or not empty_message:
            return rows
        return [
            {
                "Table_Name": "TRATAMIENTO GENERAL",
                "Message": self._extract_matching_message(soup, empty_message),
            }
        ]

    def _parse_nomenclature(self, soup):
        base_fields = {
            "Sección": self._extract_labeled_value(soup, "Sección"),
            "Capítulo:": self._extract_labeled_value(soup, "Capítulo:"),
            "Fecha inicio de vigencia:": self._extract_labeled_value(soup, "Fecha inicio de vigencia:"),
            "Fecha fin de vigencia:": self._extract_labeled_value(soup, "Fecha fin de vigencia:"),
        }
        base_fields = {key: value for key, value in base_fields.items() if value}

        rows = []
        rows.extend(self._parse_named_table_rows(soup, "Código de Mercancías", base_fields))

        additional_rows = self._parse_named_table_rows(soup, "Códigos adicionales", base_fields)
        if additional_rows:
            rows.extend(additional_rows)
        else:
            additional_message = self._extract_text_after_label(
                soup,
                "Códigos adicionales",
                self.NOMENCLATURE_BLOCK_LABELS,
            )
            if additional_message:
                rows.append(
                    {
                        **base_fields,
                        "Record_Type": "Códigos adicionales",
                        "Table_Name": "Códigos adicionales",
                        "Message": additional_message,
                    }
                )

        rows.extend(
            self._parse_named_table_rows(
                soup,
                "Unidades de medida",
                base_fields,
                selected_headers={"Código", "Descripción"},
            )
        )

        for label in (
            "Clasificadores estadísticos",
            "Descripciones mínimas",
            "Criterios de Clasificación",
        ):
            content = self._extract_text_after_label(
                soup,
                label,
                self.NOMENCLATURE_BLOCK_LABELS,
            )
            if content:
                rows.append(
                    {
                        **base_fields,
                        "Record_Type": label,
                        "Table_Name": label,
                        "Content": content,
                    }
                )

        if not rows and base_fields:
            rows.append({**base_fields, "Record_Type": "Resumen", "Table_Name": "Nomenclatura"})

        return {
            SECTION_ROWS_KEY: rows,
            SECTION_STATUS_KEY: None if rows else "No data found",
        }

    def _parse_standard_section(self, soup, empty_message=""):
        rows = self._standard_table_rows(soup, empty_message=empty_message)
        return {
            SECTION_ROWS_KEY: rows,
            SECTION_STATUS_KEY: None if rows else "No data found",
        }

    def extract_section_data(self, section_label):
        """Extract all required visible information from the currently open section."""
        soup = BeautifulSoup(self.driver.page_source, "html.parser")
        parsers = {
            "Derechos e impuestos": lambda page: self._parse_standard_section(page),
            "Nomenclatura": self._parse_nomenclature,
            "Restricciones": lambda page: self._parse_standard_section(page),
            "Cuotas": lambda page: self._parse_standard_section(
                page,
                empty_message=(
                    "No se han encontrado cuotas/contingentes "
                    "para el inciso consultado"
                ),
            ),
        }
        data = parsers.get(section_label, lambda page: {SECTION_ROWS_KEY: []})(soup)
        if data.get(SECTION_ROWS_KEY):
            return data
        return data or {SECTION_ROWS_KEY: [], SECTION_STATUS_KEY: "No data found"}

    _SECTION_JS = """
    var form = document.forms['frmBuscar'];
    if (!form) return 'no form';
    var label = arguments[0], target = arguments[1], marker = arguments[2];
    var inputs = form.getElementsByTagName('input'), button = null;
    for (var i = 0; i < inputs.length; i++) {
      if (inputs[i].type === 'submit' &&
          (inputs[i].value === label || inputs[i].value.indexOf(label) !== -1)) {
        button = inputs[i]; break;
      }
    }
    if (!button) return 'no button';
    window.open('about:blank', target);
    form.target = target;
    var hidden = document.createElement('input');
    hidden.type = 'hidden'; hidden.name = button.name; hidden.value = button.value;
    hidden.id = marker; form.appendChild(hidden); form.submit(); return 'ok';
    """

    _RESTORE_FORM_JS = """
    var form = document.forms['frmBuscar'];
    if (form) { form.target = ''; var marker = document.getElementById(arguments[0]);
      if (marker) marker.parentNode.removeChild(marker); }
    """

    def extract_section_in_new_tab(self, section_label) -> Optional[Dict]:
        """Open one SAT section in a temporary tab without losing the search session."""
        main_handle = self.driver.current_window_handle
        before = set(self.driver.window_handles)
        marker = f"__sat_section_{re.sub(r'[^A-Za-z0-9]', '_', section_label)}"
        target = f"sat_{re.sub(r'[^A-Za-z0-9]', '_', section_label)}"
        try:
            result = self.driver.execute_script(self._SECTION_JS, section_label, target, marker)
            if result != "ok":
                logger.warning("Could not open section %s: %s", section_label, result)
                return None
            time.sleep(DUTIES_TAB_DELAY if section_label == "Derechos e impuestos" else SECTION_DELAY)
            handles = set(self.driver.window_handles) - before
            if not handles:
                logger.warning("Temporary tab did not open for %s", section_label)
                return None
            self.driver.switch_to.window(handles.pop())
            data = self.extract_section_data(section_label)
            self.driver.close()
            return data
        except Exception:
            logger.exception("Error extracting section %s", section_label)
            return None
        finally:
            if main_handle in self.driver.window_handles:
                self.driver.switch_to.window(main_handle)
                self.driver.execute_script(self._RESTORE_FORM_JS, marker)

    @staticmethod
    def _prefix_section_data(section_label, data):
        prefix = re.sub(r"[^A-Za-z0-9]+", "_", section_label).strip("_")
        return {f"{prefix}__{key}": value for key, value in data.items()}

    @staticmethod
    def _safe_sheet_name(name: str) -> str:
        safe_name = re.sub(r"[\[\]:*?/\\]", "_", name).strip()
        return safe_name[:31] or "Sheet"

    @staticmethod
    def _build_section_status(overall_status: str, section_status: str) -> str:
        if section_status is not None and section_status != "":
            return section_status
        return overall_status

    @staticmethod
    def _build_overall_status(section_statuses: Dict[str, str]) -> str:
        failed_sections = [
            section_label
            for section_label, status in section_statuses.items()
            if status == "Section extraction failed"
        ]
        section_issues = [
            f"{section_label}: {status}"
            for section_label, status in section_statuses.items()
            if status not in {"Success", "Section extraction failed"}
        ]
        if failed_sections and section_issues:
            return (
                f"Missing sections: {', '.join(failed_sections)}; "
                f"Section issues: {'; '.join(section_issues)}"
            )
        if failed_sections:
            return f"Missing sections: {', '.join(failed_sections)}"
        if section_issues:
            return f"Section issues: {'; '.join(section_issues)}"
        return "Success"

    @staticmethod
    def _build_section_rows(result: Dict, section_label: str) -> List[Dict]:
        section_statuses = result.get("Section_Statuses", {})
        overall_status = result.get("Status", "")
        base_row = {
            "HS_Code": result.get("HS_Code", ""),
            "Status": SATTariffScraper._build_section_status(
                overall_status,
                section_statuses.get(section_label),
            ),
            "Overall_Status": overall_status,
        }
        section_rows = result.get("Sections", {}).get(section_label, {}).get(SECTION_ROWS_KEY, [])
        if section_rows:
            return [{**base_row, **section_row} for section_row in section_rows]
        return [base_row]

    @staticmethod
    def _format_worksheet(worksheet):
        border = Border(*(Side(style="thin"),) * 4)
        for column in worksheet.columns:
            letter = column[0].column_letter
            worksheet.column_dimensions[letter].width = min(
                max(len(str(cell.value or "")) for cell in column) + 2, 60
            )
        for cell in worksheet[1]:
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
            cell.border = border
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border

    def scrape_hs_code(self, hs_code: str, input_index: Optional[int] = None):
        result = {
            INPUT_INDEX_KEY: input_index,
            "HS_Code": hs_code,
            "Sections": {section_label: {} for section_label in SECTION_LABELS},
            "Section_Statuses": {},
        }
        if self.check_for_captcha():
            if not self.manual_captcha or not self.handle_captcha_manual():
                result["Status"] = "CAPTCHA solving failed"
                return result
        if not self.enter_hs_code(hs_code) or not self.click_search_button():
            result["Status"] = "Failed to submit HS code"
            return result
        if self.check_for_captcha() and (not self.manual_captcha or not self.handle_captcha_manual()):
            result["Status"] = "CAPTCHA solving failed after search"
            return result

        for section_label in SECTION_LABELS:
            logger.info("Extracting %s for HS %s", section_label, hs_code)
            section_data = self.extract_section_in_new_tab(section_label)
            if section_data is None:
                result["Section_Statuses"][section_label] = "Section extraction failed"
            else:
                result["Sections"][section_label] = section_data
                section_status = section_data.get(SECTION_STATUS_KEY)
                result["Section_Statuses"][section_label] = section_status or "Success"
        result["Status"] = self._build_overall_status(result["Section_Statuses"])
        return result

    @staticmethod
    def _is_successful_result(result: Dict) -> bool:
        return result.get("Status") == "Success"

    @staticmethod
    def _should_retry_result(result: Dict) -> bool:
        status = result.get("Status", "")
        return bool(status) and status != "Success" and not status.startswith("CAPTCHA")

    def _sort_results(self):
        self.results.sort(
            key=lambda result: (
                result.get(INPUT_INDEX_KEY)
                if result.get(INPUT_INDEX_KEY) is not None
                else len(self.input_order)
            )
        )

    def upsert_result(self, result: Dict):
        input_index = result.get(INPUT_INDEX_KEY)
        if input_index is not None:
            for index, existing in enumerate(self.results):
                if existing.get(INPUT_INDEX_KEY) == input_index:
                    self.results[index] = result
                    self._sort_results()
                    return
        self.results.append(result)
        self._sort_results()

    def completed_indexes(self) -> set:
        return {
            result.get(INPUT_INDEX_KEY)
            for result in self.results
            if result.get(INPUT_INDEX_KEY) is not None and self._is_successful_result(result)
        }

    @staticmethod
    def _resolve_input_index(result: Dict, saved_input_order: List[str], used_indexes: set):
        input_index = result.get(INPUT_INDEX_KEY)
        if input_index is not None:
            return input_index

        hs_code = result.get("HS_Code")
        if not hs_code:
            return None

        for index, saved_hs_code in enumerate(saved_input_order):
            if index in used_indexes:
                continue
            if saved_hs_code == hs_code:
                return index
        return None

    def save_state(self, state_file: str, output_file: str):
        path = Path(state_file)
        payload = {
            "output_file": output_file,
            "input_order": self.input_order,
            "results": self.results,
        }
        path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )

    def load_state(self, state_file: str):
        path = Path(state_file)
        if not path.exists():
            logger.info("State file not found, starting fresh: %s", state_file)
            return False

        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            logger.warning("State file could not be loaded, starting fresh: %s", exc)
            self.results = []
            self.input_order = []
            return False
        stored_order = payload.get("input_order", [])
        if stored_order and not self.input_order:
            self.input_order = stored_order
        self.results = []
        used_indexes = set()
        for result in payload.get("results", []):
            normalized_result = {**result}
            resolved_index = self._resolve_input_index(normalized_result, stored_order, used_indexes)
            if resolved_index is None:
                logger.warning(
                    "Skipping state entry without a reliable input position for HS code %s",
                    normalized_result.get("HS_Code"),
                )
                continue
            normalized_result[INPUT_INDEX_KEY] = resolved_index
            used_indexes.add(resolved_index)
            self.upsert_result(normalized_result)
        self._sort_results()
        logger.info("Loaded %d prior HS code results from %s", len(self.results), state_file)
        return True

    def save_progress(self, output_file: str, state_file: Optional[str] = None):
        self.export_to_excel(output_file)
        if state_file:
            self.save_state(state_file, output_file)

    def scrape_multiple(self, hs_codes: List[str], output_file: str, state_file: Optional[str] = None):
        self.navigate_to_portal()
        for index, hs_code in enumerate(hs_codes):
            if index in self.completed_indexes():
                logger.info("Skipping previously completed HS code %s at position %d", hs_code, index + 1)
                continue

            attempt = 0
            result = None
            while True:
                attempt += 1
                result = self.scrape_hs_code(hs_code, input_index=index)
                result.setdefault(INPUT_INDEX_KEY, index)
                if not self._should_retry_result(result) or attempt > self.max_retries:
                    break
                backoff = self.retry_backoff * (2 ** (attempt - 1))
                logger.warning(
                    "Retrying HS code %s after status %r (retry %d/%d) in %s seconds",
                    hs_code,
                    result.get("Status"),
                    attempt,
                    self.max_retries,
                    backoff,
                )
                if backoff > 0:
                    time.sleep(backoff)

            self.upsert_result(result)
            self.save_progress(output_file, state_file)
            if result["Status"].startswith("CAPTCHA"):
                break
            if index < len(hs_codes) - 1:
                time.sleep(self.delay_between_codes)
        logger.info("Completed scraping %d HS codes", len(self.results))

    def export_to_excel(self, output_file="sat_tariff_data.xlsx"):
        with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
            for section_label in SECTION_LABELS:
                rows = []
                for result in self.results:
                    rows.extend(self._build_section_rows(result, section_label))
                df = pd.DataFrame(rows)
                if df.empty:
                    df = pd.DataFrame(columns=["HS_Code", "Status", "Overall_Status"])
                ordered_columns = ["HS_Code", "Status", "Overall_Status"] + [
                    column
                    for column in df.columns
                    if column not in {"HS_Code", "Status", "Overall_Status"}
                ]
                df = df.reindex(columns=ordered_columns)
                sheet_name = self._safe_sheet_name(section_label)
                df.to_excel(writer, sheet_name=sheet_name, index=False)
                self._format_worksheet(writer.sheets[sheet_name])

    def close_browser(self):
        if self.driver:
            self.driver.quit()

    def run(
        self,
        hs_codes: List[str],
        output_file="sat_tariff_data.xlsx",
        resume=False,
        state_file=None,
        use_state_input_order=False,
    ):
        requested_hs_codes = list(hs_codes)
        self.input_order = list(requested_hs_codes)
        self.results = []
        try:
            if resume and state_file:
                if not requested_hs_codes:
                    self.input_order = []
                state_loaded = self.load_state(state_file)
                if requested_hs_codes:
                    self.input_order = list(requested_hs_codes)
                    self.results = [
                        result
                        for result in self.results
                        if result.get(INPUT_INDEX_KEY) is not None
                        and result.get(INPUT_INDEX_KEY) < len(self.input_order)
                        and result.get("HS_Code") == self.input_order[result.get(INPUT_INDEX_KEY)]
                    ]
                elif use_state_input_order and state_loaded:
                    requested_hs_codes = list(self.input_order)
                else:
                    self.input_order = []
            self._sort_results()
            if not requested_hs_codes:
                self.save_progress(output_file, state_file)
                return

            if len(self.completed_indexes()) < len(requested_hs_codes):
                self.start_browser()
                self.scrape_multiple(
                    requested_hs_codes,
                    output_file=output_file,
                    state_file=state_file,
                )
            self.save_progress(output_file, state_file)
        finally:
            self.close_browser()


if __name__ == "__main__":
    SATTariffScraper(headless=False, manual_captcha=True).run(["0101210000"])
